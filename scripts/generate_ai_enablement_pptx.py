"""
Generate the Oz Radio "AI enablement" deck series (native, editable PowerPoint).

Outputs (content/ai-enablement/):
  ai-enablement-1-beginner.pptx      12 slides  "AI, Demystified"
  ai-enablement-2-practitioner.pptx  12 slides  "Building with AI"
  ai-enablement-3-expert.pptx        12 slides  "Under the Hood"
  ai-enablement-demo-beginner.pptx    5 slides  ~3 minute public live-demo cut
  ai-enablement-demo-beginner-fr.pptx 5 slides  same demo cut in French (fr-FR live demo)

Every slide has a title (in the title placeholder, so Oz Radio's parser picks it
up), short bullets, an original diagram built only from native shapes
(rounded rectangles, ovals, connectors, freeforms), a footer with the level and
slide number, and speaker notes written to be read aloud verbatim by the avatar.

Style: white 16:9 slides, ink text, blue primary accent, restrained amber
emphasis (the "Build" presentation style), Calibri everywhere (LibreOffice in
the app container renders it with the metric-compatible Carlito font).

The output is deterministic: fixed core properties and fixed ZIP timestamps,
so regenerating unchanged content produces byte-identical files.

Run: python scripts/generate_ai_enablement_pptx.py
"""
from __future__ import annotations

import io
import math
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT_DIR = Path(__file__).resolve().parent.parent / "content" / "ai-enablement"

FIXED_TIME = datetime(2026, 10, 1, 9, 0, 0)
ZIP_TIME = (2026, 10, 1, 9, 0, 0)

FONT = "Calibri"
MONO = "Courier New"

# Build-style palette
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x16, 0x20, 0x2F)
MUTED = RGBColor(0x53, 0x61, 0x7A)
LINE = RGBColor(0xE2, 0xE8, 0xF2)
CARD_LINE = RGBColor(0xB8, 0xC4, 0xD6)
AZURE = RGBColor(0x0F, 0x6C, 0xBD)
AZURE_SOFT = RGBColor(0xEA, 0xF3, 0xFC)
AZURE_LINE = RGBColor(0x8D, 0xBB, 0xE6)
AMBER = RGBColor(0xC2, 0x62, 0x0A)
AMBER_SOFT = RGBColor(0xFF, 0xF4, 0xE4)
AMBER_LINE = RGBColor(0xE8, 0xA8, 0x55)
NODE = RGBColor(0xF7, 0xF9, 0xFC)
OK = RGBColor(0x1F, 0x9D, 0x6B)

# kind -> (fill, line, text colour)
KINDS = {
    "neutral": (PAPER, CARD_LINE, INK),
    "blue": (AZURE_SOFT, AZURE_LINE, INK),
    "solid": (AZURE, AZURE, PAPER),
    "amber": (AMBER_SOFT, AMBER_LINE, INK),
    "amber-solid": (AMBER, AMBER, PAPER),
    "node": (NODE, LINE, INK),
    "ghost": (None, CARD_LINE, MUTED),
}

SLIDE_W, SLIDE_H = 13.333333, 7.5
MARGIN = 0.55

# Diagram panel on content slides
PANEL = (5.35, 1.95, 7.43, 4.7)


# --------------------------------------------------------------------------
# Slide text (verbatim from the reviewed outline; notes are read aloud)
# --------------------------------------------------------------------------
SLIDE_TEXT = {
    '1.1': {
        "title": 'AI, Demystified',
        "bullets": [
            'A plain-language tour of AI today',
            'No coding or math required',
            'Ask questions anytime',
        ],
        "notes": (
            'Hello, and welcome. Over the next few minutes, I am going to walk you through the ideas behind '
            'modern artificial intelligence in plain language, without any coding or mathematics. By the end, '
            'you will understand what a model actually is, what people mean when they talk about an agent, '
            'and how the voice you are listening to right now works. I am, in fact, a live example of several '
            'of these ideas working together: a language model choosing words, a speech system turning text '
            'into sound, and an avatar bringing it to life on screen. Treat every slide as one small, '
            'friendly piece of a bigger picture, and feel free to keep your questions for the end.'
        ),
    },
    '1.2': {
        "title": 'What is a model?',
        "bullets": [
            'Pattern-matching built from huge examples',
            'Learns statistics, not facts directly',
            'Predicts the most likely next output',
        ],
        "notes": (
            'So what is a model, really? Think of it as a very large pattern finder. During training, it is '
            'shown enormous amounts of text or other examples, and it gradually learns statistical patterns '
            'in how those examples go together: which words tend to follow which other words, and so on. It '
            'does not store facts the way a database does. Instead, when you give it a new question, it '
            'predicts the most likely continuation based on everything it absorbed during training. That is '
            'why a model can sometimes sound confident and still be wrong, and why checking its answers '
            'against a trustworthy source remains a good habit, especially for anything that really matters.'
        ),
    },
    '1.3': {
        "title": 'Large language models vs. small language models',
        "bullets": [
            'Large models: broad knowledge, more resources',
            'Small models: compact, efficient, focused',
            'Trade-off: capability vs. cost and speed',
            'Choice depends on the task',
        ],
        "notes": (
            'Within AI, you will often hear two related terms: large language models and small language '
            'models. Both learn patterns from text, but they differ in size and the resources needed to run '
            'them. Large language models generally offer broader knowledge and stronger reasoning across '
            'varied tasks, but they need more computing power and can be slower or costlier to run. Small '
            'language models are more compact. They are cheaper to operate, can often run on a laptop or even '
            'a phone, and work well for focused tasks like summarizing one document or answering questions '
            'about a single topic. Neither is simply better; the right choice depends on what you are trying '
            'to do.'
        ),
    },
    '1.4': {
        "title": 'Reasoning models, conceptually',
        "bullets": [
            'Extra "thinking" before answering',
            'Better for multi-step problems',
            'Slower and costlier per answer',
        ],
        "notes": (
            'Some models are built, or configured, to spend extra computation thinking before they answer, '
            'rather than producing the first plausible response immediately. These are often called reasoning '
            'models. Conceptually, you can picture them working through a problem in stages internally, '
            'checking different angles before committing to a final answer, somewhat like a person pausing to '
            'work through a tricky puzzle instead of blurting out a guess. This extra step tends to help with '
            'multi-stage problems, such as planning or debugging, where getting an early step wrong would '
            'derail everything after it. The trade-off is that this extra thinking takes more time and more '
            'computing resources, so reasoning models are usually reserved for harder problems rather than '
            'everyday quick questions.'
        ),
    },
    '1.5': {
        "title": 'Prompts, tokens, and the context window',
        "bullets": [
            'Prompt: your instructions to the model',
            'Tokens: small chunks of text',
            'Context window: how much it can see',
        ],
        "notes": (
            'When you type a question, that text is called a prompt. Before the model can process it, your '
            'prompt is broken into small chunks called tokens, which might be whole words, parts of words, or '
            'punctuation. The model reads and writes one token at a time. Every model also has a context '
            'window, the maximum number of tokens it can consider at once, covering your prompt, any earlier '
            'conversation, and its own reply. If a conversation grows longer than that window, older parts '
            'have to be dropped or summarized, so the model can genuinely forget earlier details in a long '
            'chat. Understanding this helps explain why being clear and concise in a prompt usually gets '
            'better results.'
        ),
    },
    '1.6': {
        "title": 'What is an agent? Model + tools + loop',
        "bullets": [
            'Agent: a model that can act',
            'Uses tools to fetch or change things',
            'Loops: think, act, observe, repeat',
        ],
        "notes": (
            'On its own, a language model only produces text. An agent is what you get when you connect that '
            'model to tools, such as a calculator, a search engine, or a company database, and let it decide '
            'when to use them. A simple agent works in a loop: it thinks about what to do next, acts by '
            'calling a tool, observes the result, and then decides whether it is done or needs another step. '
            'This loop can repeat several times before the agent gives you a final answer. That is the key '
            'difference between a chatbot that only talks and an agent that can actually look things up or '
            'take action on your behalf.'
        ),
    },
    '1.7': {
        "title": 'Tools and MCP, explained simply',
        "bullets": [
            'Tools: specific jobs a model can call',
            'MCP: a common way to connect tools',
            'Like a universal plug for AI apps',
        ],
        "notes": (
            'Tools give a model the ability to do something beyond generating text, like searching the web, '
            'reading a file, or sending a message. For any tool to work, the AI application and the tool need '
            'to speak a common language. The Model Context Protocol, usually shortened to MCP, is an open '
            'standard that defines that shared language, so any compatible application can connect to any '
            'compatible tool without custom, one-off wiring for every pair. People sometimes describe it as '
            'being a bit like a universal plug for AI applications, the same way one cable standard lets many '
            'devices connect to many accessories. It is purely a connector; it does not decide how the AI '
            'uses what it receives.'
        ),
    },
    '1.8': {
        "title": 'Speech: speech-to-text, text-to-speech, realtime voice, avatars',
        "bullets": [
            'Speech-to-text: turns speech into words',
            'Text-to-speech: turns words into speech',
            'Realtime voice: both happen live, together',
            'Avatars: a face and voice combined',
        ],
        "notes": (
            'Several speech technologies work together to make a conversation like this one possible. '
            'Speech-to-text listens to spoken audio and converts it into written words a model can read. '
            'Text-to-speech does the reverse, turning written words back into natural-sounding spoken audio. '
            'When both happen quickly enough that a person can have a live back-and-forth conversation, that '
            'is called realtime voice. Finally, an avatar adds a visual layer, a synthesized face and body '
            'that moves and speaks in time with the generated audio, which is exactly how I am appearing to '
            'you right now. Each of these pieces can also be used on its own, for example captioning a video '
            'or reading a document aloud.'
        ),
    },
    '1.9': {
        "title": 'Open vs. closed models, and “open-weight”',
        "bullets": [
            'Closed: access via a hosted service only',
            'Open-weight: downloadable model files',
            'Open Source AI: also open data and code',
        ],
        "notes": (
            'Not every model is available the same way. A closed model is only reachable through a hosted '
            'service, where you send a request and receive a response, without ever seeing its internal '
            'files. An open-weight model is different: the organization that built it publishes the trained '
            'parameters, so anyone can download and run a copy themselves, though the training data and code '
            'may still be kept private. A smaller group of models go further and qualify as Open Source AI, '
            'meaning the data information, training code, and parameters are all shared under open terms, '
            'matching a formal definition published by the Open Source Initiative. These are different points '
            'on a spectrum of openness, not a simple yes-or-no choice.'
        ),
    },
    '1.10': {
        "title": 'Safe and responsible use',
        "bullets": [
            'Check important facts independently',
            'Watch for confident-sounding mistakes',
            'Avoid sharing sensitive data carelessly',
            'Keep a person reviewing key decisions',
        ],
        "notes": (
            'As helpful as these systems are, using them responsibly takes a little care. Models can sound '
            'completely confident while still being factually wrong, so it is worth double-checking anything '
            'important against a trustworthy source. Be thoughtful about what information you share in a '
            'prompt, especially sensitive personal or business data, since that content may be processed or '
            'stored by the service you are using. For decisions that really matter, such as medical, legal, '
            'or financial choices, keep a qualified person reviewing the output rather than acting on it '
            'automatically. Organizations such as Microsoft publish public responsible AI principles '
            'precisely to help guide these kinds of everyday judgment calls.'
        ),
    },
    '1.11': {
        "title": 'Where this fits together',
        "bullets": [
            'Models predict; agents act',
            'Tools and MCP connect the two',
            'Speech makes it conversational',
        ],
        "notes": (
            "Let's connect the pieces we have covered. A model is the prediction engine at the core. An agent "
            'wraps that model with tools and a loop so it can take action, not just talk. The Model Context '
            'Protocol is the connector that lets agents reach a growing ecosystem of tools in a standard way. '
            'Speech technology, including speech-to-text, text-to-speech, and avatars, is simply another '
            'layer on top that makes the interaction feel like a natural conversation instead of typed '
            'messages. None of these ideas are mysterious on their own; they are straightforward building '
            'blocks that combine into the AI products you may already be using every day.'
        ),
    },
    '1.12': {
        "title": "Recap and what's next",
        "bullets": [
            'You now know the core vocabulary',
            'Next deck: Building with AI (practitioner)',
            'Questions welcome anytime',
        ],
        "notes": (
            'We covered a lot of ground together: what a model is and how it predicts rather than recalls, '
            'the difference between large and small language models, what makes a reasoning model different, '
            'how prompts, tokens, and the context window shape a conversation, what an agent is and how tools '
            'and the Model Context Protocol let it act, how speech and avatars work, the spectrum from closed '
            'to open-weight to Open Source AI, and some responsible habits worth carrying forward. If you '
            'want to go deeper into actually building with these tools, the next session in this series walks '
            "through that from a practitioner's point of view. Thank you for listening, and please keep your "
            'questions coming.'
        ),
    },
    '2.1': {
        "title": 'Building with AI',
        "bullets": [
            'For engineers and technical builders',
            'From single calls to full agents',
            'Assumes basic AI familiarity',
        ],
        "notes": (
            'Welcome to the practitioner session. Where the first deck in this series introduced the '
            'vocabulary of artificial intelligence, this one is aimed at people who are actually going to '
            'build with it: choosing a model, engineering context, designing an agent harness, wiring up '
            'tools through the Model Context Protocol, adding a voice layer, and evaluating whether any of it '
            'is actually good enough to ship. I am going to assume you already know what a large language '
            'model and a token are, and we will move a little faster and more technically than the beginner '
            'session did. Treat this as a practical map of the decisions you will face turning a promising '
            'demo into something reliable enough for real users, with pointers to primary documentation so '
            'you can go deeper afterward.'
        ),
    },
    '2.2': {
        "title": 'Choosing a model: frontier vs. small vs. open-weight',
        "bullets": [
            'Frontier models: strongest general capability',
            'Small models: fast, cheap, often local',
            'Open-weight: control, customization, self-hosting',
            'Decide per task, not once',
        ],
        "notes": (
            'Model selection is rarely a one-time decision; it is usually made per task, and sometimes per '
            'step inside one workflow. Frontier models, closed and hosted by a vendor, generally offer the '
            'strongest general-purpose capability, at the cost of higher latency and price per request. Small '
            'language models trade some of that capability for speed, lower cost, and the ability to run on '
            'modest hardware or even offline. Open-weight models, where the vendor publishes downloadable '
            'parameters, sit in between: you take on hosting responsibility, but gain the ability to '
            'fine-tune, inspect, and run the model within your own environment. Illustrative, generic '
            'examples as of 2026: a frontier-class model in the GPT-4.1 family, a small model in the Phi '
            'family, an open-weight model in the Llama or Mistral families. Licences differ across all three, '
            'so check usage terms before committing.'
        ),
    },
    '2.3': {
        "title": 'Context engineering',
        "bullets": [
            'System prompt: sets persistent behavior',
            'Retrieval (RAG): fetches relevant knowledge',
            'Memory: carries facts across turns',
            'Budget: fit it all in the window',
        ],
        "notes": (
            'Context engineering is the discipline of deciding exactly what goes into the context window on '
            'every call, since that window is a hard, finite budget shared by everything the model needs to '
            'see. A system prompt anchors persistent instructions and behavior. Retrieval, commonly '
            'implemented as retrieval-augmented generation, or RAG, searches an external knowledge base and '
            'injects only the most relevant passages instead of dumping entire documents into the prompt. '
            'Memory mechanisms selectively carry forward facts from earlier turns or sessions, so the model '
            'does not need the full transcript every time. Because tokens spent on any one of these come out '
            'of the same shared budget, practitioners have to actively trim, summarize, or rank what gets '
            'included, rather than assuming more context is always better.'
        ),
    },
    '2.4': {
        "title": 'The agent harness',
        "bullets": [
            'Agent loop: plan, act, observe, repeat',
            'Tool calling: structured function invocation',
            'Memory and guardrails keep it safe',
            'Evaluation and observability close the loop',
        ],
        "notes": (
            'An agent harness is the surrounding code and infrastructure that turns a raw model into a '
            'dependable agent. At its core sits the agent loop: the model plans a next step, calls a tool, '
            'observes the result, and decides whether to continue or respond. Tool calling is the structured '
            'mechanism, typically a function definition the model fills in, that lets the model request an '
            'action in a format your code can safely execute. Guardrails constrain what the agent is allowed '
            'to do or say, catching unsafe or out-of-scope requests before or after model calls. Memory '
            'components decide what persists between turns. None of this is complete without evaluation and '
            'observability: tracing every step, scoring outputs against test cases, and watching production '
            'behavior so regressions are caught before users notice them.'
        ),
    },
    '2.5': {
        "title": 'MCP in practice: host, client, server',
        "bullets": [
            'Host: your AI application',
            'Client: one connection per server',
            'Server: exposes tools, resources, prompts',
        ],
        "notes": (
            'In the Model Context Protocol, three roles matter. The host is the AI application itself, such '
            'as an editor or a custom agent runtime. For every server it wants to use, the host creates a '
            'dedicated client, which maintains one connection and handles the protocol details for that '
            'specific server. A server is a separate program that exposes context to clients through three '
            'primitives: tools, which are callable actions; resources, which are readable data such as file '
            'contents; and prompts, which are reusable interaction templates. This separation means a single '
            'host can simultaneously connect to a local file server, an internal database server, and a '
            'remote third-party server, each through its own isolated client connection, without those '
            'servers needing to know about each other.'
        ),
    },
    '2.6': {
        "title": 'MCP in practice: transports and auth',
        "bullets": [
            'Stdio: local process, no network',
            'Streamable HTTP: remote, multi-client',
            'OAuth: authorizes remote servers',
        ],
        "notes": (
            'The Model Context Protocol defines two standard transports. Stdio runs the server as a local '
            'subprocess and exchanges messages over standard input and output, which is simple and fast but '
            'limited to one client on the same machine. Streamable HTTP is designed for remote, multi-client '
            'scenarios: the server exposes a single endpoint that accepts posted messages and can stream '
            'responses back, including through server-sent events for long-running operations. Because '
            'streamable HTTP servers are reachable over a network, the protocol recommends using OAuth, '
            'specifically a profile of OAuth 2.1, so clients obtain a proper access token before calling '
            'protected tools, rather than relying on a shared static secret. Stdio servers, by contrast, '
            'typically read credentials directly from their local environment.'
        ),
    },
    '2.7': {
        "title": 'Speech stack for practitioners',
        "bullets": [
            'Chained: speech-to-text → LLM → text-to-speech',
            'Speech-to-speech: one realtime model',
            'Latency budget drives the choice',
        ],
        "notes": (
            'There are two common architectures for building a voice agent. The chained approach runs '
            'speech-to-text, then a language model, then text-to-speech as three separate steps, which is '
            'flexible and easy to debug but adds latency at each handoff. The alternative is a '
            'speech-to-speech realtime model that handles recognition, reasoning, and voice generation inside '
            'one unified pipeline, cutting handoff latency but offering less visibility into each '
            'intermediate step. Your overall latency budget, how quickly a reply must start for the '
            'conversation to feel natural, often decides which approach fits. As one concrete, illustrative '
            'example, Azure AI Speech offers both chained services and a unified realtime option called Voice '
            'Live, which also supports avatar visuals within the same session.'
        ),
    },
    '2.8': {
        "title": 'Evaluation and cost',
        "bullets": [
            'Offline evals: test sets, rubrics',
            'Online evals: live traffic monitoring',
            'Track quality, latency, and cost together',
        ],
        "notes": (
            'Before trusting an agent with real users, you need a way to measure whether it is actually good, '
            'not just whether it runs. Offline evaluation uses a curated test set and rubrics, sometimes '
            'graded by another model, to score accuracy, relevance, or safety before anything ships. Online '
            'evaluation continues that scoring against live traffic after deployment, since real user '
            'behavior rarely matches your test cases exactly. Both should be tracked alongside latency and '
            'token cost, because a slightly less accurate but much cheaper and faster configuration is '
            'sometimes the better product decision overall. Observability tooling that traces every model '
            'call and tool invocation makes it possible to see exactly where quality, speed, or cost problems '
            'originate, rather than guessing from the final output alone.'
        ),
    },
    '2.9': {
        "title": 'Designing tools agents can use well',
        "bullets": [
            'Clear names and descriptions matter',
            'Narrow scope beats one giant tool',
            'Return structured, predictable results',
        ],
        "notes": (
            'Tool quality has an outsized effect on agent reliability, because the model chooses which tool '
            'to call largely based on its name and description. Favor narrow, single-purpose tools with clear '
            'names over one broad tool that tries to do everything, since an overloaded description makes '
            'correct selection harder for the model. Document expected inputs precisely, including types and '
            'examples, and return structured, predictable output rather than free-form text the model has to '
            're-parse. Where possible, make tools idempotent, meaning calling them twice with the same input '
            'is safe, because agent loops sometimes retry a step after a timeout or an ambiguous result. '
            'These same tools can be exposed to multiple agents through the Model Context Protocol once '
            'defined this way, making good tool design reusable across projects, not just a one-off '
            'integration.'
        ),
    },
    '2.10': {
        "title": 'Guardrails and safe defaults for builders',
        "bullets": [
            'Validate tool inputs and outputs',
            "Limit permissions to what's needed",
            'Log every decision for review',
        ],
        "notes": (
            'Guardrails are the checks that sit between what an agent decides to do and what actually '
            'happens. At minimum, validate every tool input before execution and sanity-check outputs before '
            'passing them back to the model or the user. Apply least privilege: give each tool and each agent '
            'only the permissions its job actually requires, rather than broad standing access that an error '
            'or a manipulated prompt could misuse. Keep an audit log of every decision, tool call, and '
            'outcome, since that log is what lets you diagnose an incident after the fact rather than '
            'guessing. For actions with real consequences, such as sending money or deleting data, consider '
            'requiring explicit human confirmation before the agent proceeds, at least until production '
            'evidence justifies full trust.'
        ),
    },
    '2.11': {
        "title": 'A reference architecture, end to end',
        "bullets": [
            'Voice in, text out, voice back',
            'Harness orchestrates model and tools',
            'MCP connects an ecosystem of tools',
            'Evaluation wraps the whole loop',
        ],
        "notes": (
            "Let's connect every piece into one picture. A user speaks, and a speech stack, either a chained "
            'pipeline or a realtime speech-to-speech model, turns that into something the agent harness can '
            'process. The harness runs the agent loop, calling the model, invoking tools through the Model '
            "Context Protocol when needed, and applying guardrails along the way. The model's response flows "
            'back through the speech stack as synthesized voice, optionally paired with an avatar. Wrapped '
            'around this entire loop sits evaluation and observability, continuously scoring quality, '
            'latency, and cost so you can catch regressions early. Very few production systems need every '
            'piece on day one, but understanding where each one fits makes it easier to decide what to add '
            'next as requirements grow.'
        ),
    },
    '2.12': {
        "title": 'Recap and next steps',
        "bullets": [
            'Pick models per task, not once',
            'Engineer context like a scarce resource',
            'Build the harness, then the eval loop',
            'Next deck: Under the Hood (expert)',
        ],
        "notes": (
            'We covered choosing between frontier, small, and open-weight models; engineering context '
            'deliberately rather than stuffing the window; building an agent harness with a proper loop, tool '
            'calling, memory, and guardrails; using the Model Context Protocol in practice across its roles '
            'and transports; picking a speech architecture that fits your latency budget; and evaluating '
            'quality, latency, and cost together rather than in isolation. A sensible path forward is to '
            'start small: one model, one or two well-designed tools, and a basic evaluation set, then expand '
            'the harness as real usage teaches you where it actually falls short. The next session goes one '
            'level deeper, into tokenization, attention, quantization, and the security details behind all of '
            'this.'
        ),
    },
    '3.1': {
        "title": 'Under the Hood',
        "bullets": [
            'For engineers who want internals',
            'Tokenization, attention, quantization, security',
            'Assumes practitioner-level familiarity',
        ],
        "notes": (
            'This final session goes under the hood. Where the practitioner deck covered how to build with '
            'models, agents, and the Model Context Protocol, this one examines why these systems behave the '
            'way they do at a mechanical level: how tokenization and attention create both the power and the '
            'limits of a context window, what test-time compute actually buys you in a reasoning model, how '
            'quantization and local runtimes make open-weight models practical to run outside a data center, '
            'how experienced teams structure larger agent harnesses with planners and sub-agents, the '
            'specific security risks unique to the Model Context Protocol, the latency engineering behind '
            'realtime speech, and the governance structures that keep all of this accountable. I will '
            'reference original papers and primary specifications throughout, since claims at this level '
            'deserve a traceable source.'
        ),
    },
    '3.2': {
        "title": 'Tokenization and attention',
        "bullets": [
            'Subword tokenization builds the vocabulary',
            'Self-attention relates every token pair',
            'Cost grows with sequence length',
        ],
        "notes": (
            'Modern language models are built on the transformer architecture, introduced in the 2017 paper '
            '"Attention Is All You Need," which replaced earlier recurrent designs with a mechanism called '
            'self-attention. After a tokenizer splits text into subword units, self-attention computes, for '
            'every token, how strongly it should weigh every other token in the sequence when forming its '
            'next representation. This lets the model capture long-range relationships directly, rather than '
            'passing information step by step through a chain. The practical consequence is a computational '
            'and memory cost that grows roughly with the square of sequence length, because every token pair '
            'potentially needs a comparison. That quadratic growth is the root cause of context window '
            'limits, and much of the engineering effort behind longer-context models has gone into '
            'approximating or restructuring this attention computation more efficiently.'
        ),
    },
    '3.3': {
        "title": 'Long-context trade-offs: lost in the middle',
        "bullets": [
            'Bigger window does not mean uniform recall',
            'Models favor the start and end',
            'Middle content can be effectively ignored',
        ],
        "notes": (
            'Having a large context window does not guarantee a model will use all of it equally well. A '
            'widely cited 2023 study, "Lost in the Middle," by Liu and colleagues, found that language models '
            'tend to retrieve information most reliably when it appears near the beginning or the end of a '
            'long input, with accuracy dropping noticeably for relevant facts buried in the middle of the '
            'context. This matters directly for context engineering: simply appending more retrieved passages '
            'into a prompt does not mean the model will weigh them fairly. Practical mitigations include '
            'placing the most important material near the start or end, reordering retrieved chunks by '
            'relevance rather than by original document order, and keeping retrieved context as concise as '
            'possible rather than relying on sheer volume to carry the right answer through.'
        ),
    },
    '3.4': {
        "title": 'Reasoning models and test-time compute',
        "bullets": [
            'More inference-time computation, not bigger weights',
            'Internal multi-step deliberation before output',
            'Effort is tunable per request',
        ],
        "notes": (
            'Reasoning models take a different path to better performance than simply training a larger '
            'model. Instead, they spend additional computation at inference time, working through a problem '
            'across multiple internal steps before producing a final answer, an approach generally described '
            'as test-time compute. Several platforms expose this as a tunable reasoning-effort setting, '
            'letting a developer trade latency and cost for deeper deliberation on harder requests, while '
            'leaving simple requests fast and cheap. This differs from a conventional model, which generates '
            'its response in roughly one pass per output token without an explicit internal deliberation '
            'stage. Because reasoning happens through extra generated steps, it also tends to consume more '
            'tokens overall, so cost and latency should be budgeted deliberately rather than assumed to match '
            'a non-reasoning call.'
        ),
    },
    '3.5': {
        "title": 'Open-weight models in practice: quantization and local runtimes',
        "bullets": [
            'Quantization shrinks parameter precision',
            'Trades some accuracy for size and speed',
            'Different runtimes, different goals',
        ],
        "notes": (
            'Running an open-weight model locally almost always involves quantization: reducing the numeric '
            'precision used to store each parameter, for example from a 32 or 16-bit floating-point format '
            'down to 8-bit or 4-bit integers. This shrinks the memory footprint dramatically and speeds up '
            'inference, at some cost to accuracy, though well-implemented quantization often preserves most '
            "of a model's practical quality. Open-source projects such as llama.cpp popularized running "
            'quantized open-weight models efficiently on ordinary laptops and even phones, across a range of '
            'bit-widths and hardware backends. Runtimes differ in their goals: llama.cpp and Ollama target '
            'local, single-user inference with minimal setup; vLLM targets high-throughput multi-user serving '
            'using techniques such as paged memory management for attention; and ONNX Runtime, an open, '
            'cross-platform inference engine used inside platforms such as Microsoft Foundry Local, focuses '
            'on portability across different hardware.'
        ),
    },
    '3.6': {
        "title": 'Open-weight vs. open source: licences that matter',
        "bullets": [
            'Open-weight: downloadable weights, terms vary',
            'Open Source AI: data, code, and weights open',
            "OSI's definition sets the formal bar",
            'Always read the actual licence',
        ],
        "notes": (
            'The terms open-weight and open source are often used interchangeably in casual conversation, but '
            'they describe different things, and the distinction has real licensing consequences. An '
            'open-weight model is one where the trained parameters are published for download, but the '
            'training data, training code, or usage terms may still carry significant restrictions, including '
            "limits on commercial use or redistribution. The Open Source Initiative's Open Source AI "
            'Definition sets a stricter bar: it requires sufficiently detailed information about training '
            'data, the actual code used to train and run the system, and the parameters, all made available '
            'under open terms, so a skilled person could rebuild an equivalent system. Very few widely used '
            'models currently meet that full definition. Before adopting any model family, practitioners '
            'should read its actual licence text rather than assuming a label like open or open-weight '
            'implies unrestricted use.'
        ),
    },
    '3.7': {
        "title": 'Harness design patterns: planner/executor and sub-agents',
        "bullets": [
            'Planner decomposes, executor performs steps',
            'Sub-agents isolate context per sub-task',
            'Reduces single-agent context overload',
        ],
        "notes": (
            'As tasks grow more complex, a single flat agent loop starts to struggle, because every tool '
            'definition, every intermediate result, and the full conversation history all compete for the '
            'same limited context window. A common pattern splits responsibilities: a planner model '
            'decomposes a goal into smaller sub-tasks, while one or more executor agents actually carry out '
            'each sub-task, often with their own narrower toolset and context. Sub-agents can run in isolated '
            'contexts, meaning each one only sees what is relevant to its specific piece of work, then '
            'reports a concise summary back up rather than its full internal transcript. This keeps any '
            'single context window from becoming overloaded, and it allows different sub-agents to '
            'specialize, for instance one focused on retrieval and another on code execution, which generally '
            'improves both reliability and the ability to debug failures step by step.'
        ),
    },
    '3.8': {
        "title": 'Context compaction and evals as regression tests',
        "bullets": [
            'Compaction: summarize instead of truncate',
            'Regression evals catch silent quality drops',
            'Treat prompts and tools like code',
        ],
        "notes": (
            'Two practices separate a fragile agent from a maintainable one. Context compaction periodically '
            'summarizes older parts of a long-running conversation or task instead of either keeping '
            'everything, which eventually overflows the context window, or crudely truncating it, which can '
            'silently drop something important. A good compaction strategy preserves decisions, open '
            'questions, and key facts while discarding redundant back-and-forth. Separately, regression '
            'evaluations apply the same discipline used for conventional software testing to prompts, tool '
            'definitions, and model or version changes: a fixed set of representative cases is scored '
            'automatically every time something changes, so a prompt edit or a model upgrade that quietly '
            'degrades quality on an important case gets caught before it reaches production, rather than '
            'being discovered through user complaints weeks later.'
        ),
    },
    '3.9': {
        "title": 'MCP security: auth, least privilege, consent',
        "bullets": [
            'OAuth 2.1 authorizes remote servers',
            'Scope tokens to least privilege',
            'Require explicit user consent for actions',
        ],
        "notes": (
            'Because Model Context Protocol servers can expose powerful tools, their authorization model '
            'matters. The specification defines an OAuth 2.1-based flow for remote servers reached over '
            'streamable HTTP, including protected resource metadata so a client can discover the correct '
            'authorization server automatically, rather than relying on a hardcoded shared secret. Beyond '
            'authorization, the official security guidance recommends scoping every access token to the '
            'minimum set of permissions a given client actually needs, issuing separate credentials per '
            'server rather than one broad token reused everywhere, and requiring explicit, visible user '
            "consent before a server performs a sensitive action on the user's behalf. One documented risk is "
            "the confused deputy problem, where a proxy server's static client identifier combined with a "
            'cached consent cookie can let an attacker obtain an authorization code without the user '
            "genuinely approving it, which is specifically called out in MCP's security guidance."
        ),
    },
    '3.10': {
        "title": 'Tool poisoning and prompt injection',
        "bullets": [
            'Malicious tool descriptions can mislead a model',
            'Indirect injection hides instructions in content',
            'Mitigate with validation, filtering, least privilege',
        ],
        "notes": (
            'Two closely related risks deserve special attention once an agent can call external tools. Tool '
            'poisoning happens when a malicious or compromised tool definition includes misleading names, '
            'descriptions, or embedded instructions designed to manipulate how the model behaves, exploiting '
            'the fact that the model reads tool metadata as part of its own context. Prompt injection, direct '
            'or indirect, is the broader category: direct injection comes from a user deliberately crafting '
            'adversarial input, while indirect injection hides malicious instructions inside content the '
            'agent processes later, such as a webpage, email, or document, content that is never seen or '
            'approved by a human. Mitigations recommended by groups such as the Open Web Application Security '
            'Project include constraining model behavior through strict system instructions, validating and '
            'filtering both inputs and outputs, and enforcing least-privilege access so that even a '
            'successfully manipulated agent cannot reach beyond its intended scope.'
        ),
    },
    '3.11': {
        "title": 'Streaming speech latency: barge-in and turn detection',
        "bullets": [
            'Voice activity detection flags speech boundaries',
            'Barge-in lets users interrupt the agent',
            'Auto-truncation keeps transcripts honest',
        ],
        "notes": (
            'Natural spoken conversation depends on tight latency engineering around turn-taking. Voice '
            'activity detection continuously analyzes the incoming audio stream to decide when a person has '
            'started or stopped speaking, which determines when the system should begin generating a '
            "response. Barge-in support allows a user to interrupt the agent's spoken reply mid-sentence, the "
            'way people naturally interrupt each other, rather than forcing them to wait for a full response '
            'to finish. When an interruption happens, the system needs to reconcile what was actually heard '
            'by the user with what the agent intended to say; this is sometimes implemented as '
            'auto-truncation, where the recorded transcript is cut to match only the portion of audio that '
            "actually played before the interruption, so the model's own record of the conversation stays "
            'accurate rather than including words the user never heard.'
        ),
    },
    '3.12': {
        "title": 'Governance, responsible AI, and recap',
        "bullets": [
            'Risk frameworks: NIST AI RMF and similar',
            'Responsible AI principles guide real decisions',
            'Technical depth still needs human oversight',
        ],
        "notes": (
            'All of the mechanics covered in this session, attention, quantization, planners, protocol '
            'security, and voice latency, ultimately sit inside a governance layer, and that layer deserves '
            'equal engineering attention. Voluntary frameworks such as the United States National Institute '
            "of Standards and Technology's AI Risk Management Framework organize this work around mapping "
            'risks, measuring them, and managing them continuously, rather than treating responsible AI as a '
            'one-time checklist. Vendors publish their own responsible AI principles for similar reasons, '
            'covering fairness, reliability, privacy, and accountability as concrete engineering '
            'requirements, not just marketing language. To recap this entire series: a model predicts '
            'patterns, an agent adds tools and a loop, the Model Context Protocol standardizes how tools '
            'connect, speech and avatars add a natural interface, and underneath all of it sit real '
            'trade-offs in context length, test-time compute, quantization, licensing, and security that '
            'deserve the same rigor as any other production system.'
        ),
    },
    'D.1': {
        "title": "Hi, I'm an AI presenter",
        "bullets": [
            'Built from a language model',
            'Voice generated in real time',
            'Here to explain AI, simply',
        ],
        "notes": (
            'Hello. I am a generative AI system, presenting to you live through synthesized speech and an '
            'animated avatar. Everything you are hearing right now was written by a language model moments '
            'ago and turned into audio by a separate speech system. Over the next few minutes, I want to '
            'quickly show you three ideas: what a model actually is, what makes something an agent instead of '
            'just a chatbot, and how this voice and face you are watching were built. No technical background '
            'required, just a few minutes of your attention.'
        ),
    },
    'D.2': {
        "title": 'What is a model?',
        "bullets": [
            'Trained on huge amounts of text',
            'Learns patterns, not facts',
            'Predicts the next likely word',
        ],
        "notes": (
            'A model like the one generating my words right now was trained on enormous amounts of text, '
            'learning statistical patterns in how language fits together rather than memorizing facts like a '
            'database. When you ask it something, it predicts the most likely next word, over and over, until '
            'a full answer forms. That is also why it can sound completely confident while occasionally being '
            'wrong, which is worth remembering whenever accuracy really matters.'
        ),
    },
    'D.3': {
        "title": 'What is an agent?',
        "bullets": [
            'A model that can also act',
            'Uses tools: search, calendars, files',
            'Works in a loop until done',
        ],
        "notes": (
            'A plain model only produces words. An agent is a model connected to tools, like a search engine '
            'or a calendar, that it can actually use. It works in a loop: think about the next step, act by '
            'calling a tool, look at the result, and repeat until the task is actually finished. That loop is '
            'the real difference between something that only talks and something that can get things done for '
            'you.'
        ),
    },
    'D.4': {
        "title": 'How my voice works',
        "bullets": [
            'Speech-to-text hears you',
            'Text-to-speech gives me a voice',
            'An avatar adds a moving face',
        ],
        "notes": (
            'This presentation combines a few speech technologies. Speech-to-text would convert your spoken '
            'words into text a model can read. Text-to-speech takes generated words and turns them into the '
            'natural voice you are hearing now. An avatar layer then animates a synthesized face to match '
            'that voice. When all of this happens quickly enough for a live back-and-forth, it is called '
            'realtime voice, and it is exactly what is powering me right now.'
        ),
    },
    'D.5': {
        "title": 'Thanks for watching — go deeper',
        "bullets": [
            'That was the quick version',
            'Full beginner deck: twelve topics',
            'Ask us anything, anytime',
        ],
        "notes": (
            'That was the three-minute version. The full beginner session covers twelve topics in more depth, '
            'including reasoning models, tokens and context windows, open versus closed models, and '
            'responsible use, and two further sessions go deeper still, one for practitioners who build with '
            'AI and one for experts who want the mechanics underneath. Thank you for watching, and please '
            'bring your questions, whatever your starting point with artificial intelligence happens to be.'
        ),
    },
    # French demo cut (fr-FR live demo): D.1-D.5 translated once with gpt-4.1,
    # reviewed, then hard-coded here so the output stays deterministic.
    'F.1': {
        "title": 'Bonjour, je suis un présentateur IA',
        "bullets": [
            'Construit à partir d’un modèle de langage',
            'Une voix générée en temps réel',
            'Là pour expliquer l’IA simplement',
        ],
        "notes": (
            'Bonjour. Je suis un système d’IA générative, qui vous présente en direct grâce à une voix '
            'synthétique et un avatar animé. Tout ce que vous entendez en ce moment a été écrit par un modèle '
            'de langage il y a quelques instants, puis transformé en audio par un autre système vocal. Pendant '
            'les prochaines minutes, je vais vous montrer trois idées\u00a0: ce qu’est vraiment un modèle, ce qui '
            'fait d’un agent plus qu’un simple chatbot, et comment ma voix et mon visage ont été créés. Aucun '
            'bagage technique nécessaire, juste quelques minutes de votre attention.'
        ),
    },
    'F.2': {
        "title": 'Qu’est-ce qu’un modèle\u00a0?',
        "bullets": [
            'Entraîné sur d’énormes volumes de texte',
            'Apprend des schémas, pas des faits',
            'Prédit le mot suivant le plus probable',
        ],
        "notes": (
            'Un modèle comme celui qui génère mes mots a été entraîné sur d’immenses quantités de textes, en '
            'apprenant les schémas statistiques du langage plutôt qu’en mémorisant des faits comme une base de '
            'données. Quand vous lui posez une question, il prédit le mot le plus probable, encore et encore, '
            'jusqu’à former une réponse complète. C’est aussi pour cela qu’il peut sembler très sûr de lui tout '
            'en se trompant parfois, ce qu’il faut garder en tête quand la précision compte vraiment.'
        ),
    },
    'F.3': {
        "title": 'Qu’est-ce qu’un agent\u00a0?',
        "bullets": [
            'Un modèle qui peut aussi agir',
            'Utilise des outils\u00a0: recherche, agenda, fichiers',
            'Travaille en boucle jusqu’au bout',
        ],
        "notes": (
            'Un simple modèle ne produit que des mots. Un agent, c’est un modèle connecté à des outils, comme un '
            'moteur de recherche ou un agenda, qu’il peut réellement utiliser. Il fonctionne en boucle\u00a0: il '
            'réfléchit à l’étape suivante, agit en appelant un outil, regarde le résultat, et recommence jusqu’à '
            'ce que la tâche soit vraiment terminée. Cette boucle fait toute la différence entre un système qui '
            'ne fait que parler et un système capable d’agir pour vous.'
        ),
    },
    'F.4': {
        "title": 'Comment fonctionne ma voix',
        "bullets": [
            'La reconnaissance vocale vous entend',
            'La synthèse vocale me donne une voix',
            'Un avatar ajoute un visage animé',
        ],
        "notes": (
            'Cette présentation combine plusieurs technologies vocales. La reconnaissance vocale transforme vos '
            'paroles en texte lisible par un modèle. La synthèse vocale prend les mots générés et les transforme '
            'en la voix naturelle que vous entendez. Une couche avatar anime ensuite un visage synthétique pour '
            'suivre cette voix. Quand tout cela est assez rapide pour un échange en direct, on parle de voix en '
            'temps réel, et c’est exactement ce qui me fait fonctionner en ce moment.'
        ),
    },
    'F.5': {
        "title": 'Merci de votre attention — allez plus loin',
        "bullets": [
            'C’était la version express',
            'Le deck débutant complet\u00a0: douze sujets',
            'Posez-nous vos questions, à tout moment',
        ],
        "notes": (
            'C’était la version en trois minutes. La session complète pour débutants couvre douze sujets en '
            'détail, comme les modèles de raisonnement, les tokens et les fenêtres de contexte, les modèles '
            'ouverts ou fermés, et l’usage responsable. Deux autres sessions vont encore plus loin\u00a0: une pour '
            'les praticiens qui construisent avec l’IA, et une pour les experts qui veulent comprendre les '
            'mécanismes. Merci de votre attention, et n’hésitez pas à poser vos questions, quel que soit votre '
            'niveau de départ avec l’intelligence artificielle.'
        ),
    },
}


# --------------------------------------------------------------------------
# Low-level helpers
# --------------------------------------------------------------------------
def _len(v: float) -> Emu:
    return Emu(int(round(v * 914400)))


def _no_shadow(shape) -> None:
    """Drop the theme style reference (it adds a drop shadow in LibreOffice) and any inherited effects."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    shape.shadow.inherit = False


def _style_text(run, size, *, bold=False, color=INK, font=FONT, italic=False, spacing=None):
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if spacing is not None:
        run._r.get_or_add_rPr().set("spc", str(int(spacing)))


def _frame(tf, *, margin=0.08, anchor=MSO_ANCHOR.MIDDLE, wrap=True):
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = _len(margin)
    tf.margin_top = tf.margin_bottom = _len(margin * 0.6)
    tf.vertical_anchor = anchor


def _write_lines(tf, lines, *, align=PP_ALIGN.CENTER):
    """lines: list of (text, size, bold, color, font, italic)."""
    first = True
    for text, size, bold, color, font, italic in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        run = p.add_run()
        run.text = text
        _style_text(run, size, bold=bold, color=color, font=font, italic=italic)


def text(slide, x, y, w, h, value, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, italic=False, spacing=None, margin=0.02, name=None):
    box = slide.shapes.add_textbox(_len(x), _len(y), _len(w), _len(h))
    _frame(box.text_frame, margin=margin, anchor=anchor)
    for i, line in enumerate(value.split("\n")):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        _style_text(run, size, bold=bold, color=color, font=font, italic=italic, spacing=spacing)
    if name:
        box.name = name
    return box


def shape(slide, kind_shape, x, y, w, h, *, kind="neutral", fill=None, line=None, line_w=1.25,
          dash=False, radius=0.12, rotation=0.0, name=None):
    shp = slide.shapes.add_shape(kind_shape, _len(x), _len(y), _len(w), _len(h))
    k_fill, k_line, _ = KINDS[kind]
    fill = fill if fill is not None else k_fill
    line = line if line is not None else k_line
    if kind_shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if rotation:
        shp.rotation = rotation
    _no_shadow(shp)
    if name:
        shp.name = name
    return shp


def card(slide, x, y, w, h, title="", sub=None, *, kind="neutral", size=16, sub_size=13, bold=True,
         shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         font=FONT, dash=False, line_w=1.25, radius=0.12, margin=0.08):
    shp = shape(slide, shape_type, x, y, w, h, kind=kind, dash=dash, line_w=line_w, radius=radius,
                name=(title or "card").replace("\n", " "))
    _, _, ink = KINDS[kind]
    _frame(shp.text_frame, margin=margin, anchor=anchor)
    lines = []
    for t in title.split("\n") if title else []:
        lines.append((t, size, bold, ink, font, False))
    if sub:
        sub_color = ink if kind in ("solid", "amber-solid") else MUTED
        for t in sub.split("\n"):
            lines.append((t, sub_size, False, sub_color, FONT, False))
    if lines:
        _write_lines(shp.text_frame, lines, align=align)
    return shp


def pill(slide, x, y, w, h, value, *, kind="amber", size=13, bold=True, font=FONT):
    return card(slide, x, y, w, h, value, kind=kind, size=size, bold=bold, radius=h / 2, font=font,
                margin=0.05)


def _line_ends(line_elm, *, head=False, tail=True):
    ln = line_elm
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            el = ln.makeelement(qn(tag), {"type": "triangle", "w": "med", "len": "med"})
            ln.append(el)


def arrow(slide, x1, y1, x2, y2, *, color=AZURE, width=1.75, dash=False, head=False, tail=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _len(x1), _len(y1), _len(x2), _len(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    _line_ends(c.line._get_or_add_ln(), head=head, tail=tail)
    style = c._element.find(qn("p:style"))
    if style is not None:
        c._element.remove(style)
    return c


def poly(slide, pts, *, color=AZURE, width=1.75, dash=False, tail=False, closed=False, fill=None,
         name=None):
    fb = slide.shapes.build_freeform(_len(pts[0][0]), _len(pts[0][1]), scale=1.0)
    fb.add_line_segments([(_len(px), _len(py)) for px, py in pts[1:]], close=closed)
    shp = fb.convert_to_shape()
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = color
        shp.line.width = Pt(width)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if tail:
            _line_ends(shp.line._get_or_add_ln(), tail=True)
    _no_shadow(shp)
    if name:
        shp.name = name
    return shp


def edge_point(cx, cy, r, tx, ty, pad=0.06):
    d = math.hypot(tx - cx, ty - cy) or 1.0
    return cx + (tx - cx) * (r + pad) / d, cy + (ty - cy) * (r + pad) / d


def wave(slide, x, y, w, amp, *, color=AZURE, width=2.0, cycles=2.0, steps=28):
    pts = [(x + w * i / steps, y + amp * math.sin(2 * math.pi * cycles * i / steps)) for i in range(steps + 1)]
    return poly(slide, pts, color=color, width=width)


def checkmark(slide, cx, cy, d, *, color=PAPER, width=2.5):
    pts = [(cx - d * 0.22, cy + d * 0.0), (cx - d * 0.05, cy + d * 0.17), (cx + d * 0.24, cy - d * 0.16)]
    return poly(slide, pts, color=color, width=width)


def cross(slide, cx, cy, d, *, color=PAPER, width=2.5):
    r = d * 0.18
    poly(slide, [(cx - r, cy - r), (cx + r, cy + r)], color=color, width=width)
    poly(slide, [(cx - r, cy + r), (cx + r, cy - r)], color=color, width=width)


def clock(slide, cx, cy, d, *, color=AZURE):
    shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=PAPER, line=color, line_w=1.5)
    poly(slide, [(cx, cy - d * 0.32), (cx, cy), (cx + d * 0.24, cy + d * 0.1)], color=color, width=1.5)


def gauge(slide, cx, cy, r, *, color=AZURE, value=0.7):
    pts = [(cx - r * math.cos(math.pi * i / 16), cy - r * math.sin(math.pi * i / 16)) for i in range(17)]
    poly(slide, pts, color=CARD_LINE, width=3.0)
    k = int(16 * value)
    poly(slide, pts[: k + 1], color=color, width=3.0)
    ang = math.pi * value
    poly(slide, [(cx, cy), (cx - r * 0.75 * math.cos(ang), cy - r * 0.75 * math.sin(ang))], color=INK, width=1.75)


def shield(slide, x, y, w, h, *, fill=AZURE_SOFT, line=AZURE):
    rel = [(0, 0.13), (0.5, 0), (1, 0.13), (1, 0.5), (0.9, 0.72), (0.72, 0.88), (0.5, 1),
           (0.28, 0.88), (0.1, 0.72), (0, 0.5)]
    pts = [(x + w * a, y + h * b) for a, b in rel]
    return poly(slide, pts, color=line, width=2.0, closed=True, fill=fill, name="shield")


def bubble(slide, x, y, w, h, *, wave_color=AZURE):
    b = shape(slide, MSO_SHAPE.OVAL_CALLOUT, x, y, w, h, kind="neutral", line=AZURE_LINE, line_w=1.5,
              name="chat bubble")
    wave(slide, x + w * 0.25, y + h * 0.48, w * 0.5, h * 0.11, color=wave_color, width=2.0)
    return b


def outcome(slide, value, *, y=5.85, x=5.65, w=6.83, h=0.52, kind="amber", size=15):
    return card(slide, x, y, w, h, value, kind=kind, size=size, bold=True, radius=0.1)


def flow(slide, items, *, x0, x1, y, h, gap=0.45, size=16, sub_size=13, arrow_color=AZURE):
    """items: list of (title, sub, kind). Draws cards left->right with arrows."""
    n = len(items)
    w = (x1 - x0 - gap * (n - 1)) / n
    xs = [x0 + i * (w + gap) for i in range(n)]
    for (title, sub, kind), x in zip(items, xs):
        card(slide, x, y, w, h, title, sub, kind=kind, size=size, sub_size=sub_size)
    for i in range(n - 1):
        arrow(slide, xs[i] + w + 0.04, y + h / 2, xs[i + 1] - 0.04, y + h / 2, color=arrow_color)
    return xs, w


def loop3(slide, labels, *, cx=9.07, top=2.95, bottom=5.45, half=2.05, d=1.45, center_note="repeat until done"):
    """Three circles on a triangle with clockwise arrows: top -> right -> left -> top."""
    centers = [(cx, top), (cx + half, bottom), (cx - half, bottom)]
    kinds = ["solid", "blue", "neutral"]
    for (title, sub), (ox, oy), kind in zip(labels, centers, kinds):
        card(slide, ox - d / 2, oy - d / 2, d, d, title, sub, kind=kind, size=17, sub_size=13,
             shape_type=MSO_SHAPE.OVAL, margin=0.05)
    r = d / 2
    for i in range(3):
        (ax, ay), (bx, by) = centers[i], centers[(i + 1) % 3]
        sx, sy = edge_point(ax, ay, r, bx, by)
        ex, ey = edge_point(bx, by, r, ax, ay, pad=0.1)
        arrow(slide, sx, sy, ex, ey, width=2.25)
    text(slide, cx - 1.1, (top + bottom) / 2 + 0.25, 2.2, 0.4, center_note, size=13, italic=True,
         color=MUTED, align=PP_ALIGN.CENTER)


def set_bullet(paragraph, *, indent=0.3, color=AZURE):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(_len(indent)))
    pPr.set("indent", str(-_len(indent)))
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "100000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))


def bullets(slide, x, y, w, h, items, *, size=19, space_after=14):
    box = slide.shapes.add_textbox(_len(x), _len(y), _len(w), _len(h))
    box.name = "bullets"
    _frame(box.text_frame, margin=0.02, anchor=MSO_ANCHOR.TOP)
    for i, item in enumerate(items):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = 1.05
        run = p.add_run()
        run.text = item
        _style_text(run, size, color=INK)
        set_bullet(p)
    return box


# --------------------------------------------------------------------------
# Diagrams: Deck 1 (Beginner)
# --------------------------------------------------------------------------
def v_title_beginner(s):
    card(s, 7.95, 3.3, 4.3, 1.45, "AI, Demystified", "models · agents · tools · voice", kind="solid",
         size=28, sub_size=15)
    bubble(s, 7.75, 1.75, 1.45, 0.95)
    bubble(s, 10.95, 1.95, 1.45, 0.95)
    bubble(s, 9.35, 5.25, 1.45, 0.95)


def d_model(s):
    xs, w = flow(s, [("Lots of examples", "text, images", "neutral"),
                     ("Model", "learned patterns", "solid"),
                     ("Prediction", "most likely output", "neutral")],
                 x0=5.7, x1=12.45, y=2.85, h=1.6, gap=0.5, size=17)
    for i, label in enumerate(["training", "your question"]):
        cx = xs[i] + w + 0.25
        text(s, cx - 0.85, 4.55, 1.7, 0.35, label, size=13, italic=True, color=MUTED, align=PP_ALIGN.CENTER)
    outcome(s, "Learns statistics, not a database of facts", y=5.55)


def d_scale(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.17, 5.48, 1.8, 0.14, fill=MUTED, line=None, radius=0.05)
    shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 8.62, 4.66, 0.9, 0.84, fill=MUTED, line=None)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.47, 4.56, 5.2, 0.12, fill=INK, line=None, radius=0.05,
          rotation=-7)
    card(s, 5.65, 3.08, 2.55, 1.55, "Large language model", "broad knowledge,\nmore resources", kind="solid",
         size=16)
    card(s, 10.25, 2.72, 2.2, 1.2, "Small language model", "compact, fast, focused", kind="blue", size=15)
    text(s, 10.25, 2.25, 2.2, 0.4, "lighter, rises faster", size=12, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)
    outcome(s, "Trade-off: capability vs. cost and speed", y=5.85)


def d_reasoning(s):
    card(s, 5.7, 3.3, 1.5, 0.95, "Question", kind="neutral")
    card(s, 11.0, 3.3, 1.5, 0.95, "Answer", kind="solid")
    arrow(s, 7.24, 3.55, 10.96, 3.55, width=2.25)
    pill(s, 8.3, 2.95, 1.6, 0.4, "quick model", kind="blue")
    pts = [(6.45, 4.27), (6.45, 4.9), (7.2, 5.55), (7.95, 4.9), (8.7, 5.55), (9.45, 4.9), (10.2, 5.55),
           (11.75, 5.55), (11.75, 4.32)]
    poly(s, pts, color=AMBER, width=2.25, tail=True)
    for i, (px, py) in enumerate([(7.2, 5.55), (8.7, 5.55), (10.2, 5.55)], start=1):
        card(s, px - 0.19, py - 0.19, 0.38, 0.38, str(i), kind="amber-solid", size=12,
             shape_type=MSO_SHAPE.OVAL, margin=0.0)
    pill(s, 7.35, 5.95, 3.7, 0.42, "reasoning model: extra internal steps", kind="amber")


def d_tokens(s):
    card(s, 5.7, 2.2, 6.85, 0.62, "Prompt:  \u201cSummarize this long report for me\u201d", kind="neutral",
         size=16, bold=False, align=PP_ALIGN.LEFT, margin=0.18)
    arrow(s, 7.6, 2.86, 7.6, 3.36)
    text(s, 7.75, 2.9, 2.4, 0.4, "split into tokens", size=13, italic=True, color=MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.65, 3.45, 5.75, 1.15, kind="ghost", line=AZURE, dash=True,
          line_w=1.75, name="context window")
    toks = ["Sum", "mar", "ize", "this", "long", "report", "for"]
    for i, t in enumerate(toks):
        card(s, 5.8 + i * 0.79, 3.72, 0.7, 0.6, t, kind="blue", size=14, bold=False, margin=0.02)
    card(s, 11.55, 3.95, 0.6, 0.55, "me", kind="ghost", size=14, bold=False, margin=0.0, dash=True)
    s.shapes[-1].rotation = 14
    card(s, 12.0, 4.6, 0.55, 0.5, "\u2026", kind="ghost", size=14, bold=False, margin=0.0, dash=True)
    s.shapes[-1].rotation = 32
    text(s, 5.65, 4.66, 5.75, 0.4, "Context window: the most tokens the model can see at once", size=13,
         color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 11.35, 5.15, 1.3, 0.4, "cut off", size=13, italic=True, color=AMBER, align=PP_ALIGN.CENTER)
    outcome(s, "Long chats can push earlier details out of the window", y=5.7)


def d_agent_loop(s):
    loop3(s, [("Think", None), ("Act", "use a tool"), ("Observe", "the result")])


def d_mcp_plug(s):
    card(s, 5.7, 3.68, 2.15, 1.1, "AI application", kind="solid", size=17)
    text(s, 5.6, 4.9, 2.35, 0.7, "one standard,\nmany tools", size=13, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)
    for label, y in (("Files", 2.35), ("Calendar", 3.84), ("Database", 5.33)):
        yc = y + 0.39
        arrow(s, 7.9, 4.23, 10.4, yc, color=MUTED, width=1.75, tail=False)
        mx, my = 7.9 + 0.55 * 2.5, 4.23 + 0.55 * (yc - 4.23)
        pill(s, mx - 0.37, my - 0.19, 0.74, 0.38, "MCP", kind="amber", size=12)
        card(s, 10.4, y, 2.1, 0.78, label, kind="neutral", size=16)


def d_speech(s):
    labels = [("Microphone", None, "neutral"), ("Speech\nto text", None, "blue"),
              ("Text\nto speech", None, "blue"), ("Avatar\nface", None, "solid")]
    xs, w = flow(s, labels, x0=5.7, x1=12.45, y=3.55, h=1.15, gap=0.42, size=16)
    text(s, xs[1] + w - 0.45, 4.8, 1.3, 0.35, "via a model", size=12, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)
    x_a, x_b = xs[0] + w / 2, xs[3] + w / 2
    cx, rx, ry, base = (x_a + x_b) / 2, (x_b - x_a) / 2, 1.0, 3.5
    pts = [(cx - rx * math.cos(math.pi * i / 24), base - ry * math.sin(math.pi * i / 24)) for i in range(25)]
    poly(s, pts, color=AMBER, width=2.0, dash=True, tail=True)
    pill(s, cx - 1.35, base - ry - 0.21, 2.7, 0.42, "realtime: all live, together", kind="amber")
    outcome(s, "Each piece also works alone: captions, read-aloud", y=5.6, kind="blue")


def d_openness(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.65, 2.15, 6.9, 4.3, kind="amber", radius=0.18)
    text(s, 5.9, 2.27, 6.4, 0.4, "Open Source AI  \u00b7  data information + code + weights", size=15, bold=True)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.05, 2.85, 6.1, 3.0, kind="neutral", line=AZURE_LINE, radius=0.16)
    text(s, 6.3, 2.97, 5.6, 0.4, "Open-weight  \u00b7  downloadable model files", size=15, bold=True)
    card(s, 6.45, 3.55, 5.3, 1.9, "Closed", "service access only:\nsend a request, get a response", kind="blue",
         size=17)
    text(s, 5.9, 5.93, 6.4, 0.4, "Openness grows outward: a spectrum, not yes or no", size=13, italic=True,
         color=MUTED, align=PP_ALIGN.CENTER)


def d_safety(s):
    shield(s, 8.2, 2.7, 1.75, 2.15)
    text(s, 8.2, 3.25, 1.75, 0.8, "Safe\nuse", size=17, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    for title, cx, cy, d, note in (("Verify", 6.45, 3.55, 1.3, "check key facts"),
                                   ("Privacy", 11.7, 3.55, 1.3, "share carefully"),
                                   ("Human\nreview", 9.075, 5.75, 1.2, None)):
        card(s, cx - d / 2, cy - d / 2, d, d, title, kind="blue", size=15, shape_type=MSO_SHAPE.OVAL,
             margin=0.03)
        if note:
            text(s, cx - 0.9, cy + d / 2 + 0.08, 1.8, 0.35, note, size=13, italic=True, color=MUTED,
                 align=PP_ALIGN.CENTER)
    arrow(s, 8.18, 3.55, 7.14, 3.55, color=MUTED, tail=False)
    arrow(s, 9.97, 3.55, 11.03, 3.55, color=MUTED, tail=False)
    arrow(s, 9.075, 4.87, 9.075, 5.13, color=MUTED, tail=False)


def d_stack(s):
    card(s, 5.7, 2.3, 4.7, 0.95, "Agent", "acts: model + tools + loop", kind="solid", size=17)
    card(s, 5.7, 3.4, 4.7, 0.95, "Tools via MCP", "the standard connector", kind="amber", size=17)
    card(s, 5.7, 4.5, 4.7, 0.95, "Model", "the prediction engine", kind="neutral", size=17)
    bubble(s, 10.7, 2.25, 1.85, 1.05)
    text(s, 10.6, 3.5, 2.05, 0.4, "Voice layer", size=15, bold=True, align=PP_ALIGN.CENTER)
    text(s, 10.6, 3.87, 2.05, 0.85, "speech-to-text\ntext-to-speech\navatar", size=12, color=MUTED,
         align=PP_ALIGN.CENTER)
    arrow(s, 10.45, 2.78, 10.66, 2.78, color=AZURE, tail=False, dash=True)
    outcome(s, "Simple building blocks that combine into everyday AI", y=5.75, kind="blue")


def d_checklist(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.05, 2.15, 6.0, 4.2, kind="neutral", line=AZURE_LINE, radius=0.16,
          name="checklist card")
    text(s, 6.4, 2.3, 5.3, 0.45, "You can now explain", size=16, bold=True, color=AZURE)
    rows = [("Models", "predict, not recall"), ("Agents", "model + tools + loop"),
            ("Tools / MCP", "a universal plug"), ("Speech", "voices and avatars")]
    for i, (head, tail) in enumerate(rows):
        y = 2.95 + i * 0.8
        shape(s, MSO_SHAPE.OVAL, 6.45, y, 0.52, 0.52, fill=OK, line=None)
        checkmark(s, 6.71, y + 0.27, 0.52)
        box = text(s, 7.15, y + 0.02, 4.7, 0.5, head, size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "   " + tail
        _style_text(run, 15, color=MUTED)


# --------------------------------------------------------------------------
# Diagrams: Deck 2 (Practitioner)
# --------------------------------------------------------------------------
def v_title_practitioner(s):
    y = 3.55
    poly(s, [(7.8, y), (12.3, y)], color=CARD_LINE, width=3.0)
    stops = [(8.15, "Prompt", "single call", "neutral"), (10.05, "Pipeline", "chained steps", "blue"),
             (11.95, "Agent", "loop + tools", "solid")]
    for i, (cx, head, sub, kind) in enumerate(stops):
        card(s, cx - 0.32, y - 0.32, 0.64, 0.64, str(i + 1), kind=kind, size=16, shape_type=MSO_SHAPE.OVAL,
             margin=0.0)
        text(s, cx - 0.85, y + 0.45, 1.7, 0.4, head, size=18, bold=True, align=PP_ALIGN.CENTER)
        text(s, cx - 0.85, y + 0.85, 1.7, 0.35, sub, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    for a, b in ((8.15, 10.05), (10.05, 11.95)):
        arrow(s, a + 0.4, y, b - 0.4, y, width=2.25)
    text(s, 7.8, 2.2, 4.5, 0.5, "From promising demo to reliable product", size=16, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)


def d_model_grid(s):
    gx, gy, gw, gh = 6.55, 2.2, 5.85, 3.65
    for i in range(2):
        for j in range(2):
            shape(s, MSO_SHAPE.RECTANGLE, gx + i * gw / 2, gy + j * gh / 2, gw / 2, gh / 2, fill=PAPER,
                  line=LINE, line_w=1.0)
    arrow(s, gx, gy + gh, gx + gw + 0.05, gy + gh, color=INK, width=1.75)
    arrow(s, gx, gy + gh, gx, gy - 0.05, color=INK, width=1.75)
    text(s, gx, gy + gh + 0.06, gw, 0.35, "Cost / latency  (low \u2192 high)", size=13, color=MUTED,
         align=PP_ALIGN.CENTER)
    ylab = text(s, gx - 2.3, gy + gh / 2 - 0.2, 3.65, 0.38, "Capability  (low \u2192 high)", size=13,
                color=MUTED, align=PP_ALIGN.CENTER)
    ylab.rotation = 270
    for label, cx, cy, kind, dx in (("Frontier", gx + gw * 0.8, gy + gh * 0.2, "solid", -1.68),
                                    ("Open-weight", gx + gw * 0.5, gy + gh * 0.48, "amber-solid", 0.25),
                                    ("Small", gx + gw * 0.18, gy + gh * 0.8, "blue", 0.25)):
        shape(s, MSO_SHAPE.OVAL, cx - 0.18, cy - 0.18, 0.36, 0.36, kind=kind, line_w=1.5,
              line=AZURE if kind == "blue" else None)
        text(s, cx + dx, cy - 0.22, 1.4, 0.42, label, size=16, bold=True,
             align=PP_ALIGN.RIGHT if dx < 0 else PP_ALIGN.LEFT)
    text(s, gx, gy + gh + 0.4, gw, 0.35, "Illustrative placement, not a benchmark", size=12, italic=True,
         color=MUTED, align=PP_ALIGN.CENTER)


def d_context_budget(s):
    segs = [("System\nprompt", 1.15, "solid"), ("Retrieved\ncontext (RAG)", 1.75, "blue"),
            ("Memory", 1.05, "blue"), ("Conversation\nhistory", 1.5, "neutral"),
            ("Room for\nthe answer", 1.4, "amber")]
    x, y, h = 5.7, 3.45, 1.35
    total = sum(w for _, w, _ in segs)
    for label, w, kind in segs:
        card(s, x, y, w, h, label, kind=kind, size=14, shape_type=MSO_SHAPE.RECTANGLE, margin=0.03)
        x += w
    x0, x1 = 5.7, 5.7 + total
    poly(s, [(x0, 3.25), (x0, 3.05), (x1, 3.05), (x1, 3.25)], color=INK, width=1.5)
    text(s, x0, 2.5, total, 0.45, "Context window = one shared token budget", size=16, bold=True,
         align=PP_ALIGN.CENTER)
    outcome(s, "Trim, summarize and rank: more context is not always better", y=5.4, x=5.7, w=total)


def d_harness(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.9, 2.1, 6.5, 4.38, kind="ghost", line=MUTED, dash=True,
          line_w=1.5, radius=0.16, name="observability frame")
    text(s, 6.15, 2.18, 6.0, 0.4, "Observability and evaluation", size=15, bold=True, color=AZURE)
    rows = [("Guardrails", "amber"), ("Agent loop  (plan / act / observe)", "solid"), ("Tool calling", "blue"),
            ("Memory", "neutral")]
    for i, (label, kind) in enumerate(rows):
        card(s, 6.35, 2.68 + i * 0.92, 5.6, 0.76, label, kind=kind, size=17)


def d_mcp_roles(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.65, 2.15, 3.75, 4.3, kind="blue", radius=0.16, name="host")
    text(s, 5.85, 2.25, 3.4, 0.45, "Host  \u00b7  your AI application", size=15, bold=True, color=AZURE)
    servers = ["Local file server", "Database server", "Remote server"]
    for i, srv in enumerate(servers):
        y = 2.95 + i * 1.12
        card(s, 6.0, y, 3.05, 0.78, f"Client {i + 1}", "one connection", kind="neutral", size=15, sub_size=12)
        card(s, 10.25, y, 2.3, 0.78, srv, kind="solid", size=15)
        arrow(s, 9.09, y + 0.39, 10.21, y + 0.39, head=True, tail=True, color=INK, width=1.5)
    text(s, 9.95, 6.05, 2.9, 0.4, "tools \u00b7 resources \u00b7 prompts", size=13, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)


def d_transports(s):
    for x, head, mid, foot, kind in ((5.65, "Local  \u00b7  same machine", "stdio", "Server", "neutral"),
                                     (9.2, "Remote  \u00b7  network", "Streamable HTTP", "Server", "solid")):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.15, 3.35, 4.3, kind="ghost", line=MUTED, dash=True,
              line_w=1.5, radius=0.16)
        text(s, x + 0.2, 2.25, 2.95, 0.4, head, size=15, bold=True, color=AZURE)
        card(s, x + 0.5, 2.85, 2.35, 0.8, "Host", kind="blue", size=16)
        arrow(s, x + 1.675, 3.69, x + 1.675, 5.21, width=2.25)
        card(s, x + 0.5, 5.25, 2.35, 0.8, foot, kind=kind, size=16)
        pill(s, x + 0.35, 4.08, 2.65, 0.42, mid, kind="neutral", size=13)
    pill(s, 11.0, 4.62, 1.45, 0.38, "+ OAuth token", kind="amber", size=12)


def d_speech_stack(s):
    pill(s, 5.7, 2.2, 1.3, 0.38, "Chained", kind="neutral", size=13)
    xs, w = flow(s, [("Speech-to-text", None, "blue"), ("Language model", None, "solid"),
                     ("Text-to-speech", None, "blue")], x0=5.7, x1=12.5, y=2.75, h=1.0, gap=0.5, size=15)
    for x in xs:
        clock(s, x + w - 0.1, 2.75, 0.36)
    pill(s, 5.7, 4.3, 2.0, 0.38, "Speech-to-speech", kind="neutral", size=13)
    card(s, 5.7, 4.85, 6.8, 1.0, "Unified realtime speech-to-speech model", kind="solid", size=16)
    clock(s, 12.4, 4.85, 0.28)
    text(s, 5.7, 6.0, 6.8, 0.4, "Each handoff adds latency: the latency budget drives the choice", size=13,
         italic=True, color=MUTED, align=PP_ALIGN.CENTER)


def d_eval(s):
    steps = [("Test dataset", "neutral"), ("Offline evaluation  (scores)", "blue"), ("Deploy", "amber"),
             ("Production traffic", "neutral"), ("Online evaluation & monitoring", "blue")]
    for i, (label, kind) in enumerate(steps):
        y = 2.15 + i * 0.86
        card(s, 5.7, y, 3.9, 0.6, label, kind=kind, size=15)
        if i:
            arrow(s, 7.65, y - 0.25, 7.65, y - 0.02, width=1.75)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 10.15, 2.45, 2.4, 3.75, kind="neutral", line=AZURE_LINE,
          radius=0.16, name="dashboard")
    text(s, 10.15, 2.55, 2.4, 0.4, "Dashboard", size=15, bold=True, color=AZURE, align=PP_ALIGN.CENTER)
    for i, (label, val) in enumerate((("Quality", 0.75), ("Latency", 0.45), ("Cost", 0.3))):
        y = 3.05 + i * 1.03
        gauge(s, 10.75, y + 0.5, 0.36, value=val)
        text(s, 11.2, y + 0.18, 1.3, 0.4, label, size=15, bold=True)
    arrow(s, 9.64, 5.89, 10.11, 5.6, width=1.75)


def d_tool_design(s):
    card(s, 5.7, 2.4, 2.75, 1.6, "do_everything(\n  args)", kind="neutral", size=15, bold=False, font=MONO)
    card(s, 9.75, 2.4, 2.8, 1.6, "get_order_status(\n  order_id)", kind="blue", size=15, bold=False, font=MONO)
    for cx, kind, mark in ((8.45, "amber-solid", cross), (12.55, None, checkmark)):
        shape(s, MSO_SHAPE.OVAL, cx - 0.27, 2.13, 0.54, 0.54, fill=AMBER if kind else OK, line=PAPER, line_w=2.0)
        mark(s, cx, 2.4, 0.54)
    arrow(s, 8.55, 3.2, 9.68, 3.2, width=2.25)
    text(s, 8.45, 2.75, 1.3, 0.38, "redesign", size=13, italic=True, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 5.7, 4.15, 2.75, 1.2, "vague name\none giant scope\nfree-form text", size=14, color=MUTED,
         align=PP_ALIGN.CENTER)
    text(s, 9.75, 4.15, 2.8, 1.2, "clear name, typed input\nnarrow scope\nstructured result", size=14,
         color=INK, align=PP_ALIGN.CENTER)
    outcome(s, "Well-designed tools are reusable across agents via MCP", y=5.6, kind="blue")


def d_guardrails(s):
    card(s, 5.65, 2.75, 1.9, 1.0, "Agent\ndecision", kind="blue", size=16)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.0, 2.3, 2.2, 1.9, kind="amber", radius=0.14, name="guardrail checks")
    text(s, 8.0, 2.38, 2.2, 0.4, "Guardrail checks", size=15, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    for i, lab in enumerate(("validate in / out", "least privilege", "audit log")):
        pill(s, 8.18, 2.85 + i * 0.43, 1.84, 0.36, lab, kind="neutral", size=12, bold=False)
    card(s, 10.65, 2.75, 1.9, 1.0, "Action\nexecuted", kind="solid", size=16)
    arrow(s, 7.58, 3.25, 7.96, 3.25, width=2.25)
    arrow(s, 10.24, 3.25, 10.61, 3.25, width=2.25)
    arrow(s, 9.1, 4.24, 9.1, 4.86, color=AMBER, width=2.25)
    card(s, 7.75, 4.9, 2.7, 0.85, "Blocked / flagged\nfor review", kind="amber", size=15)
    text(s, 5.65, 5.95, 6.9, 0.45, "High-impact actions: ask a person to confirm first", size=14, italic=True,
         color=MUTED, align=PP_ALIGN.CENTER)


def d_reference(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.55, 2.1, 7.1, 4.38, kind="ghost", line=MUTED, dash=True,
          line_w=1.5, radius=0.16, name="evaluation frame")
    text(s, 5.75, 2.17, 6.7, 0.4, "Evaluation & observability", size=15, bold=True, color=AZURE)
    y, h = 2.75, 1.05
    boxes = [(5.75, 1.25, "User\n(voice)", "neutral"), (7.4, 1.4, "Speech\nstack", "blue"),
             (9.2, 1.75, "Agent harness\nloop + guardrails", "solid"), (11.35, 1.1, "Model", "blue")]
    for x, w, label, kind in boxes:
        card(s, x, y, w, h, label, kind=kind, size=14, margin=0.04)
    for (x, w, _, _), (nx, _, _, _) in zip(boxes, boxes[1:]):
        arrow(s, x + w + 0.03, y + 0.35, nx - 0.03, y + 0.35, width=1.75)
    arrow(s, 11.32, y + 0.75, 10.98, y + 0.75, width=1.75)
    poly(s, [(9.6, y + h + 0.02), (9.6, 4.35), (6.375, 4.35), (6.375, y + h + 0.05)], color=AMBER, width=2.0,
         tail=True)
    text(s, 6.9, 4.4, 2.6, 0.35, "voice reply (optional avatar)", size=12, italic=True, color=AMBER,
         align=PP_ALIGN.CENTER)
    card(s, 9.55, 5.05, 2.2, 0.85, "MCP tools", kind="amber", size=15)
    arrow(s, 10.35, y + h + 0.03, 10.35, 5.01, width=1.75)
    arrow(s, 10.95, 5.01, 10.95, y + h + 0.03, width=1.75)


def d_next_expert(s):
    card(s, 5.8, 3.05, 2.4, 1.4, "You are here", "Practitioner", kind="neutral", size=17)
    shape(s, MSO_SHAPE.RIGHT_ARROW, 8.45, 3.38, 1.0, 0.75, fill=AZURE, line=None, name="next arrow")
    card(s, 9.7, 2.55, 2.85, 2.4, "Expert deck:\nUnder the Hood", "tokenization \u00b7 attention\nquantization \u00b7 security",
         kind="solid", size=19, sub_size=13)
    outcome(s, "Start small: one model, two tools, a basic eval set", y=5.55, kind="amber")


# --------------------------------------------------------------------------
# Diagrams: Deck 3 (Expert)
# --------------------------------------------------------------------------
def v_title_expert(s):
    cx, cy = 10.05, 3.95
    for d, kind in ((4.4, "ghost"), (2.9, "ghost"), (1.55, "solid")):
        shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, kind=kind, line=AZURE if kind == "ghost" else None,
              line_w=2.0, fill=PAPER if kind == "ghost" else None)
    text(s, cx - 0.75, cy - 0.25, 1.5, 0.5, "Model", size=18, bold=True, color=PAPER, align=PP_ALIGN.CENTER)
    text(s, cx - 1.0, cy - 1.38, 2.0, 0.4, "Harness", size=15, bold=True, color=AZURE, align=PP_ALIGN.CENTER)
    text(s, cx - 1.1, cy + 1.52, 2.2, 0.4, "Governance & security", size=13, bold=True, color=AMBER,
         align=PP_ALIGN.CENTER)
    for k in range(8):
        a = math.pi / 8 + k * math.pi / 4
        poly(s, [(cx + 2.32 * math.cos(a), cy + 2.32 * math.sin(a)), (cx + 2.6 * math.cos(a), cy + 2.6 * math.sin(a))],
             color=AMBER, width=2.25)


def d_attention(s):
    toks = ["The", "cat", "sat", "on", "the", "mat", "it", "slept"]
    n, c = len(toks), 0.44
    gx, gy = 6.75, 2.6
    strong = {(6, 1): 0.95, (7, 1): 0.7, (7, 6): 0.85, (2, 1): 0.75, (5, 3): 0.6, (3, 2): 0.5, (4, 5): 0.55}
    for r in range(n):
        text(s, 5.6, gy + r * c + 0.02, 1.1, c, toks[r], size=13, color=INK, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, gx + r * c - 0.15, gy - 0.4, c + 0.3, 0.35, toks[r], size=11, color=MUTED, align=PP_ALIGN.CENTER)
        for col in range(n):
            w = strong.get((r, col), 0.32 if r == col else 0.06)
            mix = lambda a, b: int(round(a + (b - a) * w))
            fill = RGBColor(mix(0xFF, 0x0F), mix(0xFF, 0x6C), mix(0xFF, 0xBD))
            shape(s, MSO_SHAPE.RECTANGLE, gx + col * c, gy + r * c, c, c, fill=fill, line=LINE, line_w=0.75)
    card(s, 10.65, 2.6, 1.95, 0.95, "8 tokens", "64 pairs", kind="neutral", size=16)
    card(s, 10.65, 3.75, 1.95, 0.95, "16 tokens", "256 pairs", kind="blue", size=16)
    card(s, 10.65, 4.9, 1.95, 0.95, "Cost grows", "with length squared", kind="amber", size=16)
    text(s, 5.6, 6.12, 5.0, 0.35, "Darker cell = stronger attention (illustrative)", size=12, italic=True,
         color=MUTED, align=PP_ALIGN.CENTER)


def d_lost_middle(s):
    x0, x1, top, low = 6.05, 12.35, 2.75, 4.45
    pts = []
    for i in range(31):
        t = i / 30
        u = (2 * t - 1) ** 2
        pts.append((x0 + (x1 - x0) * t, low - (low - top) * (u ** 0.7)))
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.15, 3.75, 2.1, 1.0, kind="ghost", line=AMBER, dash=True, line_w=1.5)
    poly(s, pts, color=AZURE, width=3.0)
    text(s, 5.65, 2.15, 4.0, 0.4, "Retrieval accuracy by position", size=15, bold=True, color=AZURE)
    text(s, 8.0, 3.3, 2.4, 0.4, "middle is recalled less", size=13, italic=True, color=AMBER, align=PP_ALIGN.CENTER)
    y = 5.0
    for x, w, label, kind in ((6.05, 1.9, "Start", "blue"), (7.95, 2.5, "Middle", "neutral"),
                              (10.45, 1.9, "End", "blue")):
        card(s, x, y, w, 0.62, label, kind=kind, size=15, shape_type=MSO_SHAPE.RECTANGLE)
    text(s, 6.05, 5.66, 6.3, 0.35, "Long input / context window", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 5.65, 6.05, 7.0, 0.35, "Illustrative shape, after Liu et al. (2023)", size=12, italic=True,
         color=MUTED, align=PP_ALIGN.CENTER)


def d_test_time(s):
    base = 5.55
    for i, (label, hgt, d) in enumerate((("low effort", 0.9, 0.34), ("medium effort", 1.75, 0.44),
                                         ("high effort", 2.6, 0.54))):
        x = 5.95 + i * 1.35
        card(s, x, base - hgt, 0.95, hgt, "", kind="solid" if i == 2 else "blue", shape_type=MSO_SHAPE.RECTANGLE)
        clock(s, x + 0.475, base - hgt - 0.12 - d / 2, d)
        text(s, x - 0.25, base + 0.08, 1.45, 0.35, label, size=13, color=INK, align=PP_ALIGN.CENTER)
    arrow(s, 9.85, 4.2, 10.4, 4.2, width=2.25)
    card(s, 10.45, 3.65, 2.1, 1.1, "Answer\nquality", kind="solid", size=17)
    text(s, 10.3, 4.85, 2.4, 0.8, "more tokens, more\nlatency, more cost", size=13, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)


def d_quantization(s):
    base = 4.3
    x = 5.9
    for label, d in (("32-bit", 1.45), ("16-bit", 1.15), ("8-bit", 0.9), ("4-bit", 0.7)):
        shape(s, MSO_SHAPE.CUBE, x, base - d, d, d, kind="blue", line=AZURE, line_w=1.5, name=f"{label} cube")
        text(s, x - 0.2, base + 0.05, d + 0.4, 0.35, label, size=14, bold=True, align=PP_ALIGN.CENTER)
        x += d + 0.45
    arrow(s, 5.9, 4.9, 12.2, 4.9, color=AMBER, width=2.25)
    text(s, 5.9, 4.95, 6.3, 0.35, "smaller, faster, slightly less precise", size=13, italic=True, color=AMBER,
         align=PP_ALIGN.CENTER)
    for i, (head, sub) in enumerate((("llama.cpp \u00b7 Ollama", "local, single user"),
                                     ("vLLM", "high-throughput serving"),
                                     ("ONNX Runtime", "portable across hardware"))):
        card(s, 5.7 + i * 2.32, 5.45, 2.18, 0.9, head, sub, kind="neutral", size=14, sub_size=12, margin=0.04)


def d_licence_pyramid(s):
    apex_x, top, bottom, half = 7.6, 2.3, 6.3, 2.0
    tiers = [(top, 3.75, AMBER_SOFT, AMBER_LINE), (3.85, 5.05, AZURE_SOFT, AZURE_LINE), (5.15, bottom, NODE, CARD_LINE)]
    hw = lambda y: half * (y - top) / (bottom - top)
    labels = [("Open Source AI", "+ full data information\n(OSI definition)"),
              ("+ Code", "training and inference\ncode available"),
              ("Weights", "downloadable files\n(open-weight)")]
    for i, ((y0, y1, fill, line), (head, sub)) in enumerate(zip(tiers, labels)):
        if i == 0:
            pts = [(apex_x, y0), (apex_x + hw(y1), y1), (apex_x - hw(y1), y1)]
        else:
            pts = [(apex_x - hw(y0), y0), (apex_x + hw(y0), y0), (apex_x + hw(y1), y1), (apex_x - hw(y1), y1)]
        poly(s, pts, color=line, width=1.5, closed=True, fill=fill, name=f"tier {3 - i}")
        mid = (y0 + y1) / 2 + (0.25 if i == 0 else 0)
        text(s, apex_x - 0.4, mid - 0.25, 0.8, 0.5, str(3 - i), size=18, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, color=AMBER if i == 0 else AZURE)
        arrow(s, apex_x + hw(mid) + 0.1, mid, 9.85, mid, color=CARD_LINE, width=1.25, tail=False)
        text(s, 9.95, mid - 0.45, 2.65, 0.4, head, size=16, bold=True, color=AMBER if i == 0 else INK)
        text(s, 9.95, mid - 0.07, 2.65, 0.7, sub, size=13, color=MUTED)


def d_planner(s):
    card(s, 8.0, 2.15, 2.15, 0.8, "Planner", "decomposes the goal", kind="solid", size=16, sub_size=12)
    subs = [("Sub-agent A", "retrieval"), ("Sub-agent B", "code execution"), ("Sub-agent C", "drafting")]
    for i, (head, sub) in enumerate(subs):
        x = 5.7 + i * 2.33
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.35, 2.15, 1.38, kind="ghost", line=MUTED, dash=True, line_w=1.25)
        text(s, x, 3.38, 2.15, 0.3, "isolated context", size=11, italic=True, color=MUTED, align=PP_ALIGN.CENTER)
        card(s, x + 0.15, 3.7, 1.85, 0.88, head, sub, kind="blue", size=15, sub_size=12)
        arrow(s, 9.075, 2.98, x + 1.075, 3.31, width=1.5)
        arrow(s, x + 1.075, 4.77, 9.075, 5.25, width=1.5, color=AMBER)
    card(s, 7.75, 5.3, 2.65, 0.85, "Executor / aggregator", "concise summaries back", kind="amber", size=15,
         sub_size=12)


def d_compaction(s):
    text(s, 5.65, 2.2, 1.1, 0.5, "Before", size=14, bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(8):
        card(s, 6.7 + i * 0.73, 2.2, 0.63, 0.5, f"T{i + 1}", kind="neutral", size=13, bold=False, margin=0.0)
    shape(s, MSO_SHAPE.GEAR_6, 7.95, 2.85, 0.6, 0.6, fill=AMBER, line=None, name="compaction gear")
    arrow(s, 8.25, 3.48, 8.25, 3.8, color=AMBER, width=2.0)
    text(s, 8.65, 2.92, 3.9, 0.45, "summarize older turns, keep decisions", size=13, italic=True, color=AMBER)
    text(s, 5.65, 3.88, 1.1, 0.5, "After", size=14, bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    card(s, 6.7, 3.95, 2.85, 0.38, "Summary of T1 \u2013 T5", kind="amber", size=13, bold=True, margin=0.02)
    for i in range(5, 8):
        card(s, 6.7 + i * 0.73, 3.88, 0.63, 0.5, f"T{i + 1}", kind="neutral", size=13, bold=False, margin=0.0)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.65, 4.75, 6.95, 1.7, kind="blue", radius=0.14)
    text(s, 5.85, 4.82, 6.6, 0.4, "Regression evals: treat prompts and tools like code", size=15, bold=True,
         color=AZURE)
    xs, w = flow(s, [("Prompt or model change", None, "neutral"), ("Score fixed cases", None, "neutral"),
                     ("Ship or block", None, "solid")], x0=5.85, x1=12.4, y=5.35, h=0.8, gap=0.4, size=14)


def d_mcp_security(s):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.55, 3.15, 1.05, 1.35, fill=None, line=INK, line_w=6.0, radius=0.5,
          name="padlock shackle")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.3, 3.85, 1.55, 1.3, fill=AZURE, line=None, radius=0.14,
          name="padlock body")
    shape(s, MSO_SHAPE.OVAL, 8.95, 4.25, 0.25, 0.25, fill=PAPER, line=None)
    shape(s, MSO_SHAPE.RECTANGLE, 9.035, 4.42, 0.08, 0.33, fill=PAPER, line=None)
    pill(s, 7.75, 2.2, 2.65, 0.42, "OAuth 2.1 for remote servers", kind="amber", size=13)
    arrow(s, 9.075, 2.65, 9.075, 3.08, color=AMBER, width=2.0)
    card(s, 5.7, 3.7, 1.95, 1.0, "Token scope", "least privilege", kind="blue", size=15, sub_size=12)
    card(s, 10.5, 3.7, 2.05, 1.0, "User consent", "explicit, visible", kind="blue", size=15, sub_size=12)
    card(s, 7.95, 5.55, 2.25, 0.85, "Per-server isolation", "separate credentials", kind="blue", size=15,
         sub_size=12)
    arrow(s, 7.69, 4.2, 8.26, 4.4, width=2.0)
    arrow(s, 10.46, 4.2, 9.89, 4.4, width=2.0)
    arrow(s, 9.075, 5.52, 9.075, 5.19, width=2.0)


def d_injection(s):
    shape(s, MSO_SHAPE.FOLDED_CORNER, 5.75, 2.35, 2.3, 3.0, kind="neutral", line=CARD_LINE, name="document")
    for i in range(7):
        if i in (3, 4):
            continue
        shape(s, MSO_SHAPE.RECTANGLE, 5.98, 2.65 + i * 0.32, 1.8 - (0.4 if i % 3 == 2 else 0), 0.08, fill=LINE,
              line=None)
    card(s, 5.92, 3.6, 1.95, 0.55, "hidden instruction", kind="amber", size=13, margin=0.03)
    pill(s, 5.6, 5.5, 2.6, 0.42, "also: poisoned tool descriptions", kind="amber", size=12, bold=False)
    arrow(s, 8.1, 3.85, 8.5, 3.85, color=AMBER, width=2.25)
    shield(s, 8.55, 2.95, 1.4, 1.75)
    text(s, 8.15, 4.8, 2.2, 0.6, "input / output\nfiltering", size=13, bold=True, color=AZURE, align=PP_ALIGN.CENTER)
    arrow(s, 10.0, 3.85, 10.85, 3.85, width=2.25, dash=True)
    card(s, 10.9, 3.35, 1.65, 1.0, "Agent", kind="solid", size=17)
    text(s, 10.55, 4.45, 2.2, 0.8, "least privilege limits\nthe blast radius", size=12, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)


def d_barge_in(s):
    cut = 9.55
    for label, y in (("User", 2.7), ("Agent", 4.0)):
        text(s, 5.65, y - 0.3, 1.1, 0.6, label, size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(40):
        x = 6.75 + i * 0.145
        amp = 0.12 + 0.42 * abs(math.sin(i * 0.9) * math.cos(i * 0.37))
        if x > cut + 0.05:
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.7 - amp / 2, 0.08, amp, fill=AMBER, line=None, radius=0.04)
        if x < cut - 0.1:
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.0 - amp / 2, 0.08, amp, fill=AZURE, line=None, radius=0.04)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cut + 0.05, 3.7, 12.55 - cut - 0.05, 0.6, kind="ghost", line=AZURE_LINE,
          dash=True, line_w=1.25)
    text(s, cut + 0.1, 3.8, 12.45 - cut - 0.1, 0.4, "truncated response", size=13, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER)
    poly(s, [(cut, 2.15), (cut, 4.6)], color=INK, width=1.75, dash=True)
    text(s, cut - 1.9, 2.05, 1.8, 0.6, "user starts\nspeaking", size=12, italic=True, color=INK, align=PP_ALIGN.RIGHT)
    pill(s, 6.45, 4.7, 2.35, 0.4, "voice activity detection", kind="neutral", size=12)
    pill(s, cut - 0.6, 4.7, 1.2, 0.4, "barge-in", kind="amber-solid", size=12)
    outcome(s, "Auto-truncation: the transcript keeps only audio that played", y=5.55, kind="blue")


def d_governance(s):
    cx, cy, r = 9.075, 4.3, 1.95
    shape(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, kind="ghost", line=AZURE_LINE, line_w=10.0,
          fill=PAPER, name="governance ring")
    card(s, cx - 0.85, cy - 0.85, 1.7, 1.7, "AI system", kind="solid", size=17, shape_type=MSO_SHAPE.OVAL)
    for label, dx, dy, kind in (("Govern", 0, -1, "amber-solid"), ("Map", 1, 0, "neutral"),
                                ("Measure", 0, 1, "neutral"), ("Manage", -1, 0, "neutral")):
        w = 1.45
        pill(s, cx + dx * r - w / 2, cy + dy * r - 0.25, w, 0.5, label, kind=kind, size=15)
    text(s, 5.6, 2.2, 2.4, 0.5, "NIST AI RMF\nfunctions", size=13, bold=True, color=MUTED)


# --------------------------------------------------------------------------
# Diagrams: Demo cut
# --------------------------------------------------------------------------
# Diagram labels for the demo cut, per language (EN is the reviewed original).
DEMO_LABELS = {
    "en": {
        "D.1": ("language model", "realtime voice", "avatar"),
        "D.2": ("Training\ntext", "Model", "Prediction", "next likely word",
                "Patterns, not facts: it can sound sure and still be wrong"),
        "D.3": ("Think", "Act", "call a tool", "Observe", "look at result", "repeat until done"),
        "D.4": ("Your\nvoice", "Speech\nto text", "Model", "Text to\nspeech", "Avatar",
                "Fast enough for live back-and-forth = realtime voice"),
        "D.5": ("You are here", "3-minute demo", "Full AI enablement series",
                "3 decks:\nBeginner \u00b7 Practitioner \u00b7 Expert",
                "Bring your questions, whatever your starting point"),
    },
    # Translated once with gpt-4.1, then reviewed and shortened to fit the boxes.
    "fr": {
        "D.1": ("mod\u00e8le de langage", "voix temps r\u00e9el", "avatar"),
        "D.2": ("Textes\nd\u2019entra\u00eenement", "Mod\u00e8le", "Pr\u00e9diction", "mot suivant probable",
                "Des sch\u00e9mas, pas des faits\u00a0: il peut sembler s\u00fbr et se tromper"),
        "D.3": ("R\u00e9fl\u00e9chir", "Agir", "appeler un outil", "Observer", "lire le r\u00e9sultat",
                "r\u00e9p\u00e9ter jusqu\u2019au bout"),
        "D.4": ("Votre\nvoix", "Voix en\ntexte", "Mod\u00e8le", "Texte en\nvoix", "Avatar",
                "Assez rapide pour un \u00e9change en direct = voix temps r\u00e9el"),
        "D.5": ("Vous \u00eates ici", "d\u00e9mo de 3 minutes", "S\u00e9rie compl\u00e8te AI enablement",
                "3 decks\u00a0:\nD\u00e9butant \u00b7 Praticien \u00b7 Expert",
                "Venez avec vos questions, quel que soit votre niveau"),
    },
}


def v_title_demo(s, lang="en"):
    shape(s, MSO_SHAPE.OVAL, 8.0, 1.85, 2.8, 2.8, kind="blue", name="avatar")
    shape(s, MSO_SHAPE.OVAL, 8.95, 2.3, 0.9, 0.9, fill=AZURE, line=None)
    shape(s, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 8.55, 3.35, 1.7, 1.05, fill=AZURE, line=None)
    bubble(s, 10.95, 1.75, 1.55, 1.0)
    size = 13 if lang == "en" else 12
    for i, label in enumerate(DEMO_LABELS[lang]["D.1"]):
        pill(s, 7.65 + i * 1.65, 5.15, 1.55, 0.45, label, kind="blue" if i < 2 else "solid", size=size)


def d_demo_model(s, lang="en"):
    training, model, prediction, prediction_sub, takeaway = DEMO_LABELS[lang]["D.2"]
    card(s, 5.7, 3.05, 1.75, 1.2, training, kind="neutral", size=16 if lang == "en" else 14)
    shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 7.55, 2.95, 1.4, 1.0, fill=AZURE_SOFT, line=AZURE_LINE, rotation=90,
          name="funnel")
    card(s, 9.05, 2.85, 1.6, 1.6, model, kind="solid", size=17, shape_type=MSO_SHAPE.OVAL)
    arrow(s, 10.7, 3.65, 11.0, 3.65, width=2.25)
    card(s, 11.05, 3.05, 1.5, 1.2, prediction, prediction_sub, kind="neutral", size=15, sub_size=12)
    outcome(s, takeaway, y=5.3)


def d_demo_agent(s, lang="en"):
    think, act, act_sub, observe, observe_sub, note = DEMO_LABELS[lang]["D.3"]
    loop3(s, [(think, None), (act, act_sub), (observe, observe_sub)], center_note=note)


def d_demo_voice(s, lang="en"):
    voice, stt, model, tts, avatar, takeaway = DEMO_LABELS[lang]["D.4"]
    flow(s, [(voice, None, "neutral"), (stt, None, "blue"), (model, None, "solid"),
             (tts, None, "blue"), (avatar, None, "solid")],
         x0=5.7, x1=12.5, y=3.3, h=1.15, gap=0.35, size=15 if lang == "en" else 14)
    outcome(s, takeaway, y=5.2, kind="amber")


def d_demo_next(s, lang="en"):
    here, here_sub, series, series_sub, takeaway = DEMO_LABELS[lang]["D.5"]
    card(s, 5.8, 3.05, 2.3, 1.4, here, here_sub, kind="neutral", size=17)
    shape(s, MSO_SHAPE.RIGHT_ARROW, 8.35, 3.38, 1.0, 0.75, fill=AZURE, line=None, name="next arrow")
    card(s, 9.6, 2.5, 2.95, 2.5, series, series_sub, kind="solid", size=19, sub_size=13)
    outcome(s, takeaway, y=5.55, kind="amber")


def _fr(fn):
    return lambda s: fn(s, lang="fr")


DIAGRAMS = {
    "1.1": v_title_beginner, "1.2": d_model, "1.3": d_scale, "1.4": d_reasoning, "1.5": d_tokens,
    "1.6": d_agent_loop, "1.7": d_mcp_plug, "1.8": d_speech, "1.9": d_openness, "1.10": d_safety,
    "1.11": d_stack, "1.12": d_checklist,
    "2.1": v_title_practitioner, "2.2": d_model_grid, "2.3": d_context_budget, "2.4": d_harness,
    "2.5": d_mcp_roles, "2.6": d_transports, "2.7": d_speech_stack, "2.8": d_eval, "2.9": d_tool_design,
    "2.10": d_guardrails, "2.11": d_reference, "2.12": d_next_expert,
    "3.1": v_title_expert, "3.2": d_attention, "3.3": d_lost_middle, "3.4": d_test_time, "3.5": d_quantization,
    "3.6": d_licence_pyramid, "3.7": d_planner, "3.8": d_compaction, "3.9": d_mcp_security,
    "3.10": d_injection, "3.11": d_barge_in, "3.12": d_governance,
    "D.1": v_title_demo, "D.2": d_demo_model, "D.3": d_demo_agent, "D.4": d_demo_voice, "D.5": d_demo_next,
    "F.1": _fr(v_title_demo), "F.2": _fr(d_demo_model), "F.3": _fr(d_demo_agent), "F.4": _fr(d_demo_voice),
    "F.5": _fr(d_demo_next),
}


# --------------------------------------------------------------------------
# Deck definitions
# --------------------------------------------------------------------------
DECKS = [
    {
        "file": "ai-enablement-1-beginner.pptx",
        "level": "Beginner",
        "title": "AI enablement 1 \u2014 Beginner: AI, Demystified",
        "series": "Oz Radio \u00b7 AI enablement series \u00b7 Part 1 of 3",
        "audience": "Beginner \u00b7 no prior AI knowledge",
        "tagline": "Models, agents, tools and voice in plain language",
        "eyebrows": ["Welcome", "Foundations", "Models", "Models", "Foundations", "Agents", "Agents", "Speech",
                     "Models", "Responsible AI", "Big picture", "Recap"],
        "ids": [f"1.{i}" for i in range(1, 13)],
    },
    {
        "file": "ai-enablement-2-practitioner.pptx",
        "level": "Practitioner",
        "title": "AI enablement 2 \u2014 Practitioner: Building with AI",
        "series": "Oz Radio \u00b7 AI enablement series \u00b7 Part 2 of 3",
        "audience": "Practitioner \u00b7 builds with AI tools",
        "tagline": "From prototype to production",
        "eyebrows": ["Welcome", "Model choice", "Context", "Agents", "MCP", "MCP", "Speech", "Evaluation",
                     "Tools", "Safety", "Architecture", "Recap"],
        "ids": [f"2.{i}" for i in range(1, 13)],
    },
    {
        "file": "ai-enablement-3-expert.pptx",
        "level": "Expert",
        "title": "AI enablement 3 \u2014 Expert: Under the Hood",
        "series": "Oz Radio \u00b7 AI enablement series \u00b7 Part 3 of 3",
        "audience": "Expert \u00b7 wants the internals",
        "tagline": "Mechanics, trade-offs and security",
        "eyebrows": ["Welcome", "Transformers", "Long context", "Reasoning", "Local inference", "Licensing",
                     "Harness patterns", "Reliability", "MCP security", "Security", "Realtime speech",
                     "Governance & recap"],
        "ids": [f"3.{i}" for i in range(1, 13)],
    },
    {
        "file": "ai-enablement-demo-beginner.pptx",
        "level": "Beginner demo",
        "title": "AI enablement \u2014 Beginner demo cut (about 3 minutes)",
        "series": "Oz Radio \u00b7 AI enablement \u00b7 Live demo",
        "audience": "Beginner \u00b7 three-minute demo",
        "tagline": "A quick, live look at how AI works",
        "eyebrows": ["Live demo", "Models", "Agents", "Speech", "Go deeper"],
        "ids": [f"D.{i}" for i in range(1, 6)],
    },
    {
        "file": "ai-enablement-demo-beginner-fr.pptx",
        "level": "Démo débutant",
        "title": "AI enablement \u2014 Démo débutant (environ 3 minutes)",
        "series": "Oz Radio \u00b7 AI enablement \u00b7 Démo en direct",
        "audience": "Débutant \u00b7 démo de trois minutes",
        "tagline": "Un aperçu rapide de l’IA, en direct",
        "eyebrows": ["Démo en direct", "Modèles", "Agents", "Voix", "Pour aller plus loin"],
        "ids": [f"F.{i}" for i in range(1, 6)],
    },
]


# --------------------------------------------------------------------------
# Slide builders
# --------------------------------------------------------------------------
def _set_title(slide, value, *, x, y, w, h, size):
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = _len(x), _len(y), _len(w), _len(h)
    tf = title.text_frame
    _frame(tf, margin=0.02, anchor=MSO_ANCHOR.TOP)
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 0.95
    run = p.add_run()
    run.text = value
    _style_text(run, size, bold=True, color=INK)
    return title


def _footer(slide, deck, n, total):
    poly(slide, [(MARGIN, 6.86), (SLIDE_W - MARGIN, 6.86)], color=LINE, width=1.0)
    text(slide, MARGIN, 6.92, 8.0, 0.35, f"Oz Radio \u00b7 AI enablement \u00b7 {deck['level']}", size=11,
         color=MUTED, name="footer")
    text(slide, SLIDE_W - MARGIN - 2.0, 6.92, 2.0, 0.35, f"{n} / {total}", size=11, color=MUTED,
         align=PP_ALIGN.RIGHT, name="slide number")


def _eyebrow(slide, value, *, y=0.42):
    text(slide, MARGIN, y, 10.0, 0.32, value.upper(), size=12, bold=True, color=AZURE, spacing=150,
         name="eyebrow")


def _title_slide(slide, deck, data, sid):
    _eyebrow(slide, deck["series"], y=0.5)
    title_text = data["title"]
    size = 46 if len(title_text) <= 18 else 40
    _set_title(slide, title_text, x=MARGIN, y=1.15, w=6.5, h=1.55, size=size)
    text(slide, MARGIN, 2.72, 6.5, 0.5, deck["tagline"], size=20, color=AZURE)
    pill(slide, MARGIN, 3.35, 4.0, 0.44, deck["audience"], kind="amber", size=14)
    bullets(slide, MARGIN, 4.1, 6.3, 2.5, data["bullets"], size=20, space_after=12)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.35, 1.2, 5.43, 5.4, kind="node", radius=0.18, name="visual panel")
    DIAGRAMS[sid](slide)


def _content_slide(slide, deck, data, sid, eyebrow):
    _eyebrow(slide, f"{deck['level']} \u00b7 {eyebrow}")
    title_text = data["title"]
    size = 30 if len(title_text) <= 52 else 27
    _set_title(slide, title_text, x=MARGIN, y=0.74, w=SLIDE_W - 2 * MARGIN, h=1.1, size=size)
    bullets(slide, MARGIN, 2.1, 4.5, 4.5, data["bullets"])
    x, y, w, h = PANEL
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, kind="node", radius=0.18, name="diagram panel")
    DIAGRAMS[sid](slide)


def build_deck(deck: dict) -> bytes:
    prs = Presentation()
    prs.slide_width = _len(SLIDE_W)
    prs.slide_height = _len(SLIDE_H)
    layout = prs.slide_layouts[5]  # "Title Only": gives each slide a real title placeholder
    total = len(deck["ids"])
    for n, (sid, eyebrow) in enumerate(zip(deck["ids"], deck["eyebrows"]), start=1):
        data = SLIDE_TEXT[sid]
        slide = prs.slides.add_slide(layout)
        if n == 1:
            _title_slide(slide, deck, data, sid)
        else:
            _content_slide(slide, deck, data, sid, eyebrow)
        _footer(slide, deck, n, total)
        slide.notes_slide.notes_text_frame.text = data["notes"]

    cp = prs.core_properties
    cp.title = deck["title"]
    cp.subject = "Oz Radio AI enablement series"
    cp.author = "Oz Radio"
    cp.last_modified_by = "Oz Radio"
    cp.keywords = "AI enablement; Oz Radio; " + deck["level"]
    cp.comments = "Generated by scripts/generate_ai_enablement_pptx.py"
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.last_printed = FIXED_TIME
    cp.revision = 1

    raw = io.BytesIO()
    prs.save(raw)
    return _normalize_zip(raw.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    """Rewrite the package with fixed entry timestamps so output is byte-stable."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, src.read(info.filename))
    return out.getvalue()


def deck_word_count(deck: dict) -> int:
    return sum(len(SLIDE_TEXT[sid]["notes"].split()) for sid in deck["ids"])


def main(argv: list[str]) -> int:
    out_dir = Path(argv[1]) if len(argv) > 1 else OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    for deck in DECKS:
        path = out_dir / deck["file"]
        path.write_bytes(build_deck(deck))
        words = deck_word_count(deck)
        print(f"Created {path} ({len(deck['ids'])} slides, {words} words of notes, "
              f"~{words / 150:.1f} min at 150 wpm)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
