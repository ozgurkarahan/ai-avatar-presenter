"""Checks for the committed AI enablement deck series (content/ai-enablement/).

Opens each PPTX with python-pptx only (no LibreOffice) and verifies slide
counts, titles, speaker notes and narration length per level. The decks are
also compared with a fresh in-memory build from
scripts/generate_ai_enablement_pptx.py so stale outputs are caught.
"""
from __future__ import annotations

import importlib.util
import io
from pathlib import Path

import pytest
from pptx import Presentation

REPO = Path(__file__).resolve().parents[3]
DECK_DIR = REPO / "content" / "ai-enablement"
GENERATOR = REPO / "scripts" / "generate_ai_enablement_pptx.py"

# file -> (slide count, inclusive word band per slide's notes)
EXPECTED = {
    "ai-enablement-1-beginner.pptx": (12, (95, 130)),
    "ai-enablement-2-practitioner.pptx": (12, (105, 150)),
    "ai-enablement-3-expert.pptx": (12, (115, 160)),
    "ai-enablement-demo-beginner.pptx": (5, (55, 100)),
    "ai-enablement-demo-beginner-fr.pptx": (5, (55, 100)),
}
DEMO_FILE = "ai-enablement-demo-beginner.pptx"
DEMO_FR_FILE = "ai-enablement-demo-beginner-fr.pptx"
DEMO_MAX_WORDS = 420  # ~3 minutes at 150 wpm, within the 180 s public demo session

FORBIDDEN = ("api_key", "apikey", "password", "secret=", "sig=", "accountkey")


def _load_generator():
    spec = importlib.util.spec_from_file_location("generate_ai_enablement_pptx", GENERATOR)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _slides(path: Path):
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None else ""
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        body = "\n".join(
            s.text_frame.text for s in slide.shapes if s.has_text_frame and s != slide.shapes.title
        )
        out.append({"title": title, "notes": notes, "body": body})
    return prs, out


@pytest.fixture(scope="module")
def generator():
    return _load_generator()


@pytest.mark.parametrize("filename", sorted(EXPECTED))
def test_deck_structure(filename):
    path = DECK_DIR / filename
    assert path.is_file(), f"missing {path}; run python scripts/generate_ai_enablement_pptx.py"
    count, (lo, hi) = EXPECTED[filename]
    prs, slides = _slides(path)

    assert len(slides) == count
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    for i, s in enumerate(slides, start=1):
        assert s["title"], f"{filename} slide {i}: empty title"
        assert s["notes"], f"{filename} slide {i}: empty speaker notes"
        words = len(s["notes"].split())
        assert lo <= words <= hi, f"{filename} slide {i}: {words} words outside {lo}-{hi}"
        assert "Oz Radio \u00b7 AI enablement \u00b7" in s["body"], f"{filename} slide {i}: footer missing"
        assert f"{i} / {count}" in s["body"], f"{filename} slide {i}: slide number missing"
        for needle in ("http://", "https://", "(words)"):
            assert needle not in s["notes"], f"{filename} slide {i}: notes contain {needle!r}"


@pytest.mark.parametrize("filename", [DEMO_FILE, DEMO_FR_FILE])
def test_demo_cut_fits_three_minutes(filename):
    _, slides = _slides(DECK_DIR / filename)
    total = sum(len(s["notes"].split()) for s in slides)
    assert total <= DEMO_MAX_WORDS, f"demo cut notes are {total} words (> {DEMO_MAX_WORDS})"


@pytest.mark.parametrize("filename", sorted(EXPECTED))
def test_no_secrets(filename):
    _, slides = _slides(DECK_DIR / filename)
    blob = "\n".join(s["title"] + "\n" + s["body"] + "\n" + s["notes"] for s in slides).lower()
    for needle in FORBIDDEN:
        assert needle not in blob, f"{filename} contains {needle!r}"


def test_committed_decks_match_generator(generator):
    by_file = {d["file"]: d for d in generator.DECKS}
    assert set(by_file) == set(EXPECTED)
    for filename, deck in by_file.items():
        _, committed = _slides(DECK_DIR / filename)
        fresh_prs = Presentation(io.BytesIO(generator.build_deck(deck)))
        fresh = [
            (s.shapes.title.text_frame.text.strip(), s.notes_slide.notes_text_frame.text.strip())
            for s in fresh_prs.slides
        ]
        assert [(s["title"], s["notes"]) for s in committed] == fresh, (
            f"{filename} is stale; run python scripts/generate_ai_enablement_pptx.py"
        )
        assert [s["title"] for s in committed] == [generator.SLIDE_TEXT[sid]["title"] for sid in deck["ids"]]


def test_french_demo_cut_mirrors_english(generator):
    """Same 5-slide structure as the EN demo cut, but fully translated."""
    en_prs, en = _slides(DECK_DIR / DEMO_FILE)
    fr_prs, fr = _slides(DECK_DIR / DEMO_FR_FILE)
    assert len(fr) == len(en) == 5
    assert (fr_prs.slide_width, fr_prs.slide_height) == (en_prs.slide_width, en_prs.slide_height)
    for i, (e, f) in enumerate(zip(en, fr), start=1):
        assert f["title"] != e["title"], f"slide {i}: title not translated"
        assert f["notes"] != e["notes"], f"slide {i}: notes not translated"
        assert len(f["body"]) > 0
    fr_text = " ".join(s["title"] + " " + s["notes"] for s in fr)
    for word in (" vous ", "modèle", "agent", "voix"):
        assert word in fr_text, f"French demo cut lacks {word!r}"
    en_ids = next(d for d in generator.DECKS if d["file"] == DEMO_FILE)["ids"]
    fr_ids = next(d for d in generator.DECKS if d["file"] == DEMO_FR_FILE)["ids"]
    for e_id, f_id in zip(en_ids, fr_ids):
        assert len(generator.SLIDE_TEXT[f_id]["bullets"]) == len(generator.SLIDE_TEXT[e_id]["bullets"])
    # The diagrams use translated labels too (no leftover English diagram text).
    fr_lines = {line.strip() for s in fr for line in s["body"].splitlines()}
    for leftover in generator.DEMO_LABELS["en"]["D.2"] + generator.DEMO_LABELS["en"]["D.3"]:
        if leftover in generator.DEMO_LABELS["fr"]["D.2"] + generator.DEMO_LABELS["fr"]["D.3"]:
            continue  # identical in both languages
        assert leftover.split("\n")[0] not in fr_lines, f"English label {leftover!r} in French deck"
