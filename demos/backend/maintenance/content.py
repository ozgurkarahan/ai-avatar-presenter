"""Oz Radio content maintenance CLI (operator job, runs inside the app container).

Seeds the AI enablement catalog, renders podcasts/videos for it and purges
retired content, reusing the routers' own code paths (catalog upload,
publish, translate-notes, live demo, playlists, podcast/video render workers
and library deletes) with the app's managed identity. Nothing goes through
HTTP or Easy Auth, so no principal is created or impersonated; the per-user
daily request quota (an HTTP-layer limit) is not consumed by these jobs.

Usage (from the backend directory, ``/app`` in the image)::

    python -m maintenance.content inventory [--json out.json]
    python -m maintenance.content seed [--dry-run]
    python -m maintenance.content render-podcast {beginner,practitioner,expert,demo-fr,all} [--dry-run]
    python -m maintenance.content render-video {demo-en,demo-fr,practitioner,all} [--dry-run]
    python -m maintenance.content purge --deck-title T ... --podcast-id job-... --video-id sv-... \\
        --record deleted.json [--allow-missing video:practitioner] [--dry-run]
    python -m maintenance.content retitle [--dry-run]
    python -m maintenance.content rebrand-decks [--deck ai-enablement-1-beginner ...] [--dry-run]
    python -m maintenance.content --background /tmp/all.log all [--skip video:practitioner]
    python -m maintenance.content logtail /tmp/all.log -n 40

``--background LOG`` detaches the job (new session) so a short
``az containerapp exec`` call can start a long render; ``logtail`` prints the
log with URL query strings redacted.

Every command is idempotent: decks are matched by title, the playlist by name
and podcasts/videos by library title (an existing item is only published).
Titles match either the display title (``DISPLAY_TITLES``) or the original
filename slug, so content seeded before ``retitle`` is never duplicated.
Output never contains signed URLs.

``render-podcast``/``render-video --force`` render a fresh item even when one
with the same title exists (e.g. after a rebrand); retire the previous item
afterwards with ``purge --podcast-id/--video-id`` (media only: purge touches
playlists only when ``--deck-title`` deletes decks).

``rebrand-decks`` re-renders the AI decks from the regenerated .pptx and
overwrites the slide images, slide text and search chunks of the *same*
deck ids in place (speaker notes must be identical; titles, translated
notes, published/live-demo flags and playlists are kept).
"""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import logging
import shutil
import sys
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

log = logging.getLogger("maintenance.content")

# Decks ship in the image at maintenance/decks/ (Dockerfile); locally they are
# read from the repository's content/ai-enablement/.
DECK_DIRS = (
    Path(__file__).resolve().parent / "decks",
    BACKEND_DIR.parent.parent / "content" / "ai-enablement",
)

INTERVIEWER = ("Sofia", "photo_female", "female")
EXPERT = ("Lucas", "photo_male", "male")

# language -> gender -> HD voice (must be in services.voice_catalog)
VOICES = {
    "en-US": {"female": "en-US-Ava:DragonHDLatestNeural", "male": "en-US-Andrew:DragonHDLatestNeural"},
    "fr-FR": {"female": "fr-FR-Vivienne:DragonHDLatestNeural", "male": "fr-FR-Remy:DragonHDLatestNeural"},
}


# Filename slug -> human-readable display title (decks, podcasts and videos).
DISPLAY_TITLES: dict[str, str] = {
    "ai-enablement-1-beginner": "AI, demystified · Beginner",
    "ai-enablement-2-practitioner": "Building with AI · Practitioner",
    "ai-enablement-3-expert": "Under the hood · Expert",
    "ai-enablement-demo-beginner": "AI in three minutes · Live demo",
    "ai-enablement-demo-beginner-fr": "L’IA en trois minutes · Démo en direct",
}


def display_title(slug: str) -> str:
    return DISPLAY_TITLES.get(slug, slug)


class _Titled:
    """Mixin: ``slug`` (file stem), ``title`` (display) and both as ``titles``."""

    _source: str

    @property
    def slug(self) -> str:
        return Path(self._source).stem

    @property
    def title(self) -> str:
        return display_title(self.slug)

    @property
    def titles(self) -> frozenset[str]:
        return frozenset({self.slug, self.title})


@dataclass(frozen=True)
class DeckSeed(_Titled):
    file: str
    language: str
    live_demo: bool = False
    translate_to: tuple[str, ...] = ()

    @property
    def _source(self) -> str:  # type: ignore[override]
        return self.file


SEED_DECKS: tuple[DeckSeed, ...] = (
    DeckSeed("ai-enablement-1-beginner.pptx", "en-US", translate_to=("fr-FR",)),
    DeckSeed("ai-enablement-2-practitioner.pptx", "en-US", translate_to=("fr-FR",)),
    DeckSeed("ai-enablement-3-expert.pptx", "en-US", translate_to=("fr-FR",)),
    DeckSeed("ai-enablement-demo-beginner.pptx", "en-US", live_demo=True, translate_to=("fr-FR",)),
    DeckSeed("ai-enablement-demo-beginner-fr.pptx", "fr-FR", live_demo=True, translate_to=("en-US",)),
)

PLAYLIST_TITLE = "AI enablement path"
PLAYLIST_DESCRIPTION = (
    "From first principles to production internals: AI, Demystified (Beginner), "
    "Building with AI (Practitioner), then Under the Hood (Expert)."
)
PLAYLIST_STEPS = (
    "ai-enablement-1-beginner",
    "ai-enablement-2-practitioner",
    "ai-enablement-3-expert",
)


@dataclass(frozen=True)
class PodcastSpec(_Titled):
    deck: str
    language: str
    style: str
    audience_level: str
    length: str = "medium"
    num_turns: int = 8

    @property
    def _source(self) -> str:  # type: ignore[override]
        return self.deck


PODCASTS: dict[str, PodcastSpec] = {
    "beginner": PodcastSpec("ai-enablement-1-beginner.pptx", "en-US", "explainer", "beginner"),
    "practitioner": PodcastSpec("ai-enablement-2-practitioner.pptx", "en-US", "casual", "practitioner"),
    "expert": PodcastSpec("ai-enablement-3-expert.pptx", "en-US", "debate", "expert"),
    "demo-fr": PodcastSpec("ai-enablement-demo-beginner-fr.pptx", "fr-FR", "casual", "beginner"),
}


@dataclass(frozen=True)
class VideoSpec(_Titled):
    deck: str
    language: str
    avatar: str
    gender: str
    audience_level: str
    style: str = "explainer"

    @property
    def _source(self) -> str:  # type: ignore[override]
        return self.deck

    @property
    def voice(self) -> str:
        return VOICES[self.language][self.gender]


VIDEOS: dict[str, VideoSpec] = {
    "demo-en": VideoSpec("ai-enablement-demo-beginner.pptx", "en-US", "photo_female", "female", "beginner"),
    "demo-fr": VideoSpec("ai-enablement-demo-beginner-fr.pptx", "fr-FR", "photo_male", "male", "beginner"),
    "practitioner": VideoSpec("ai-enablement-2-practitioner.pptx", "en-US", "photo_female", "female",
                              "practitioner"),
}


class MaintenanceError(RuntimeError):
    pass


def _say(msg: str) -> None:
    print(f"[{datetime.now(timezone.utc).strftime('%H:%M:%S')}] {msg}", flush=True)


def _cfg():
    from config import load_config

    return load_config()


def _deck_path(name: str, deck_dir: Optional[Path] = None) -> Path:
    dirs = (Path(deck_dir),) if deck_dir else DECK_DIRS
    for d in dirs:
        p = d / name
        if p.is_file():
            return p
    raise MaintenanceError(f"deck file {name} not found in {[str(d) for d in dirs]}")


def _staged_copy(src: Path, kind: str) -> Path:
    """Copy a deck under data/uploads/<kind>/ like the upload routes do."""
    tmp_dir = Path("data/uploads") / kind
    tmp_dir.mkdir(parents=True, exist_ok=True)
    dst = tmp_dir / f"{uuid.uuid4().hex}{src.suffix.lower()}"
    shutil.copyfile(src, dst)
    return dst


def _strip_urls(item: dict) -> dict:
    return {k: v for k, v in item.items() if not k.endswith("_url")}


# ---------------------------------------------------------------------------
# Inventory
# ---------------------------------------------------------------------------
def inventory(cfg) -> dict[str, Any]:
    from routers import decks as dr
    from routers import podcast as pr
    from routers import video as vr

    out: dict[str, Any] = {"captured_at": datetime.now(timezone.utc).isoformat()}
    store = dr._store(cfg)
    if not store or not store.cosmos_available:
        raise MaintenanceError("Cosmos DB unavailable")
    out["decks"] = [dr._doc_to_summary(d).model_dump() for d in store.list_decks()
                    if isinstance(d, dict) and d.get("id")]
    out["playlists"] = [
        {"id": p.get("id"), "title": p.get("title"), "status": p.get("status", "active"),
         "steps": [{"deck_id": s.get("deck_id"), "deck_title": s.get("deck_title"), "order": s.get("order")}
                   for s in sorted(p.get("steps") or [], key=lambda s: s.get("order", 0))]}
        for p in store.list_playlists()
    ]
    plib, vlib = pr._get_library(cfg), vr._get_library(cfg)
    out["podcasts"] = [_strip_urls(i) for i in plib.list()] if plib.available else []
    out["videos"] = [_strip_urls(i) for i in vlib.list()] if vlib.available else []
    return out


# ---------------------------------------------------------------------------
# Seed decks + playlist
# ---------------------------------------------------------------------------
async def seed(cfg, *, dry_run: bool = False, deck_dir: Optional[Path] = None) -> dict[str, Any]:
    from routers import decks as dr
    from routers import playlists as plr

    store = dr._store(cfg)
    if not store or not store.cosmos_available:
        raise MaintenanceError("Cosmos DB unavailable")
    by_title: dict[str, dict] = {}
    for d in store.list_decks():
        if isinstance(d, dict) and d.get("id"):
            by_title.setdefault(dr._doc_to_summary(d).title, d)

    report: dict[str, Any] = {"decks": [], "playlist": None, "dry_run": dry_run}
    ids: dict[str, Optional[str]] = {}
    for spec in SEED_DECKS:
        path = _deck_path(spec.file, deck_dir)
        entry: dict[str, Any] = {"title": spec.title, "language": spec.language}
        doc = by_title.get(spec.title) or by_title.get(spec.slug)
        if doc:
            deck_id = doc["id"]
            entry["action"] = "exists"
            if (doc.get("language") or "en-US") != spec.language:
                entry["warning"] = f"existing deck language is {doc.get('language')}"
        elif dry_run:
            deck_id = None
            entry["action"] = "would upload"
        else:
            content = path.read_bytes()
            if content[:4] != b"PK\x03\x04":
                raise MaintenanceError(f"{path} is not a .pptx")
            _say(f"upload {spec.file} ({spec.language})")
            resp = await asyncio.to_thread(dr.ingest_deck_bytes, content, spec.file, spec.language, cfg,
                                           title=spec.title)
            deck_id = resp.deck_id
            entry.update(action="uploaded", slide_count=resp.slide_count, indexed_slides=resp.indexed_slides)
        entry["deck_id"] = deck_id
        ids[spec.slug] = deck_id

        if deck_id and not dry_run:
            current = store.get_presentation(deck_id) or {}
            if current.get("published") is not True:
                dr.publish_deck(deck_id, dr.PublishRequest(published=True), cfg=cfg)
                entry["published"] = "now"
            else:
                entry["published"] = "already"
            for lang in spec.translate_to:
                tr = await dr.translate_deck_notes_for(deck_id, lang, cfg)
                entry.setdefault("translated", {})[lang] = "cached" if tr.cached else f"{len(tr.translated_notes)} slides"
            if spec.live_demo:
                current = store.get_presentation(deck_id) or {}
                if current.get("live_demo") is not True:
                    dr.set_live_demo(deck_id, dr.LiveDemoRequest(enabled=True), cfg=cfg)
                    entry["live_demo"] = "now"
                else:
                    entry["live_demo"] = "already"
        _say(f"deck {spec.title}: {entry['action']}")
        report["decks"].append(entry)

    existing = [p for p in store.list_playlists()
                if p.get("title") == PLAYLIST_TITLE and p.get("status", "active") != "broken"]
    if existing:
        report["playlist"] = {"id": existing[0].get("id"), "action": "exists"}
    elif dry_run or any(not ids.get(t) for t in PLAYLIST_STEPS):
        report["playlist"] = {"action": "would create", "steps": list(PLAYLIST_STEPS)}
    else:
        body = plr.PlaylistCreate(
            title=PLAYLIST_TITLE,
            description=PLAYLIST_DESCRIPTION,
            steps=[plr.PlaylistStepIn(deck_id=ids[t], order=i) for i, t in enumerate(PLAYLIST_STEPS)],
        )
        created = plr.create_playlist(body, cfg=cfg)
        report["playlist"] = {"id": created.id, "action": "created",
                              "steps": [s.deck_title for s in created.steps]}
    _say(f"playlist {PLAYLIST_TITLE!r}: {report['playlist']['action']}")
    return report


# ---------------------------------------------------------------------------
# Podcasts
# ---------------------------------------------------------------------------
def _podcast_roles(language: str):
    from services import avatar_registry
    from services.podcast_models import RenderRoles, RoleConfig
    from services.voice_catalog import is_allowed_voice

    roles = RenderRoles(
        interviewer=RoleConfig(display_name=INTERVIEWER[0], avatar=INTERVIEWER[1],
                               voice=VOICES[language][INTERVIEWER[2]]),
        expert=RoleConfig(display_name=EXPERT[0], avatar=EXPERT[1], voice=VOICES[language][EXPERT[2]]),
    )
    for role in (roles.interviewer, roles.expert):
        if not avatar_registry.is_known(role.avatar) or avatar_registry.get(role.avatar).id != role.avatar:
            raise MaintenanceError(f"avatar {role.avatar!r} is not registered (CUSTOM_AVATARS_JSON)")
        if not is_allowed_voice(role.voice):
            raise MaintenanceError(f"voice {role.voice!r} is not allowed")
    return roles


def _existing(lib, titles: frozenset[str] | set[str]) -> Optional[dict]:
    items = [i for i in lib.list() if i.get("title") in titles]
    return items[0] if items else None


async def render_podcast(cfg, name: str, *, dry_run: bool = False, deck_dir: Optional[Path] = None,
                         force: bool = False) -> dict:
    from routers import podcast as pr
    from services.podcast_ingest import ingest_file
    from services.podcast_models import (
        JobOutputs, JobProgress, JobState, PodcastJob, Script, ScriptRequest,
    )
    from services.podcast_script import generate_script_stream

    spec = PODCASTS[name]
    path = _deck_path(spec.deck, deck_dir)
    roles = _podcast_roles(spec.language)
    lib = pr._get_library(cfg)
    if not lib.available:
        raise MaintenanceError("podcast library storage unavailable")
    plan = {"name": name, "title": spec.title, "language": spec.language, "style": spec.style,
            "audience_level": spec.audience_level,
            "roles": {"interviewer": roles.interviewer.model_dump(), "expert": roles.expert.model_dump()}}
    found = _existing(lib, spec.titles)
    if found and force:
        plan["replaces"] = found["job_id"]
    elif found:
        plan.update(action="exists", job_id=found["job_id"], published=found.get("published"))
        if not found.get("published") and not dry_run:
            lib.set_published(found["job_id"], True)
            plan["published"] = True
        _say(f"podcast {name}: exists ({found['job_id']})")
        return plan
    if dry_run:
        plan["action"] = "would render"
        return plan

    _say(f"podcast {name}: ingest {spec.deck}")
    staged = _staged_copy(path, "podcast")
    doc = await asyncio.to_thread(ingest_file, staged, cfg.libreoffice_path or "soffice")
    doc.title = spec.title
    pr.DOCUMENTS[doc.id] = doc

    _say(f"podcast {name}: script ({spec.language}, {spec.style}, {spec.audience_level})")
    req = ScriptRequest(document_id=doc.id, language=spec.language, style=spec.style, length=spec.length,
                        num_turns=spec.num_turns, audience_level=spec.audience_level)
    turns = [t async for t in generate_script_stream(cfg, doc, req)]
    if not turns:
        raise MaintenanceError("script generation returned no turns")
    pr._check_turns(turns)
    script = Script(id=f"scr-{uuid.uuid4().hex[:8]}", document_id=doc.id, language=spec.language,
                    style=spec.style, length=spec.length, audience_level=spec.audience_level, turns=turns)
    pr.SCRIPTS[script.id] = script

    now = pr._now()
    job = PodcastJob(id=f"job-{uuid.uuid4().hex[:8]}", script_id=script.id, roles=roles,
                     layout="split_screen_with_slides", music=True, intro=True, state=JobState.queued,
                     progress=JobProgress(stage="queued", total=len(turns)), outputs=JobOutputs(),
                     created_at=now, updated_at=now)
    pr.JOBS[job.id] = job
    _say(f"podcast {name}: render {job.id} ({len(turns)} turns)")
    await pr._render_job_body(job.id, cfg)
    if job.state != JobState.done or job.archive_state != "published":
        raise MaintenanceError(f"podcast {job.id} failed: state={job.state.value} archive={job.archive_state}")
    if not lib.set_published(job.id, True):
        raise MaintenanceError(f"podcast {job.id} rendered but publishing failed")
    plan.update(action="rendered", job_id=job.id, published=True, turns=len(turns))
    _say(f"podcast {name}: published {job.id}")
    return plan


# ---------------------------------------------------------------------------
# Videos
# ---------------------------------------------------------------------------
async def render_video(cfg, name: str, *, dry_run: bool = False, deck_dir: Optional[Path] = None,
                       force: bool = False) -> dict:
    from routers import video as vr
    from security import MAX_NARRATION_CHARS, MAX_SLIDES_PER_RENDER
    from services import avatar_registry
    from services.video_ingest import ingest
    from services.video_models import (
        JobOutputs, JobProgress, JobState, ScriptRequest, VideoJob, VideoScript,
    )
    from services.video_script import generate_script_stream
    from services.voice_catalog import is_allowed_voice

    spec = VIDEOS[name]
    path = _deck_path(spec.deck, deck_dir)
    if not is_allowed_voice(spec.voice):
        raise MaintenanceError(f"voice {spec.voice!r} is not allowed")
    # The video worker picks the presenter from the voice (gender match); make
    # sure that resolves to the requested photo avatar before spending a render.
    resolved = avatar_registry.for_voice(spec.voice).id
    if resolved != spec.avatar:
        raise MaintenanceError(f"voice {spec.voice} resolves to avatar {resolved!r}, expected {spec.avatar!r}")
    lib = vr._get_library(cfg)
    if not lib.available:
        raise MaintenanceError("video library storage unavailable")
    plan = {"name": name, "title": spec.title, "language": spec.language, "voice": spec.voice,
            "avatar": spec.avatar, "audience_level": spec.audience_level, "style": spec.style}
    found = _existing(lib, spec.titles)
    if found and force:
        plan["replaces"] = found["job_id"]
    elif found:
        plan.update(action="exists", job_id=found["job_id"], published=found.get("published"))
        if not found.get("published") and not dry_run:
            lib.set_published(found["job_id"], True)
            plan["published"] = True
        _say(f"video {name}: exists ({found['job_id']})")
        return plan
    if dry_run:
        plan["action"] = "would render"
        return plan

    _say(f"video {name}: ingest {spec.deck}")
    staged = _staged_copy(path, "video")
    doc = await asyncio.to_thread(ingest, str(staged), spec.deck)
    doc.title = spec.title
    vr.DOCUMENTS[doc.doc_id] = doc

    _say(f"video {name}: script ({spec.language}, {spec.audience_level})")
    req = ScriptRequest(language=spec.language, style=spec.style, audience_level=spec.audience_level,
                        voice=spec.voice)
    narrations = [n async for n in generate_script_stream(cfg, doc, req)]
    if not narrations:
        raise MaintenanceError("script generation returned no narrations")
    if len(narrations) > MAX_SLIDES_PER_RENDER:
        raise MaintenanceError(f"too many slides (max {MAX_SLIDES_PER_RENDER})")
    for n in narrations:
        if len(n.narration or "") > MAX_NARRATION_CHARS or (n.voice and not is_allowed_voice(n.voice)):
            raise MaintenanceError(f"slide {n.slide_index}: narration too long or voice not allowed")
    vr.SCRIPTS[doc.doc_id] = VideoScript(
        doc_id=doc.doc_id, language=spec.language, style=spec.style, audience_level=spec.audience_level,
        voice=spec.voice, narrations=sorted(narrations, key=lambda n: n.slide_index),
    )

    now = vr._now()
    job = VideoJob(job_id=f"sv-{uuid.uuid4().hex[:8]}", doc_id=doc.doc_id, state=JobState.queued,
                   progress=JobProgress(stage="queued", total=len(narrations)), outputs=JobOutputs(),
                   created_at=now, updated_at=now)
    vr.JOBS[job.job_id] = job
    _say(f"video {name}: render {job.job_id} ({len(narrations)} slides)")
    await vr._render_job_body(job.job_id, cfg)
    if job.state != JobState.done or job.archive_state != "published":
        raise MaintenanceError(f"video {job.job_id} failed: state={job.state.value} archive={job.archive_state}")
    if not lib.set_published(job.job_id, True):
        raise MaintenanceError(f"video {job.job_id} rendered but publishing failed")
    plan.update(action="rendered", job_id=job.job_id, published=True, slides=len(narrations))
    _say(f"video {name}: published {job.job_id}")
    return plan


# ---------------------------------------------------------------------------
# Purge
# ---------------------------------------------------------------------------
def missing_new_content(inv: dict, allow_missing: set[str]) -> list[str]:
    """What of the new AI catalog is not yet present + published (empty = ready)."""
    missing: list[str] = []
    decks = {d["title"]: d for d in inv["decks"]}
    for spec in SEED_DECKS:
        d = decks.get(spec.title) or decks.get(spec.slug)
        if not d or not d.get("published") or (spec.live_demo and not d.get("live_demo")):
            missing.append(f"deck:{spec.title}")
    if not any(p.get("title") == PLAYLIST_TITLE and p.get("status") != "broken" for p in inv["playlists"]):
        missing.append(f"playlist:{PLAYLIST_TITLE}")
    for kind, specs, items in (("podcast", PODCASTS, inv["podcasts"]), ("video", VIDEOS, inv["videos"])):
        published = {i.get("title") for i in items if i.get("published")}
        for name, spec in specs.items():
            key = f"{kind}:{name}"
            if not (spec.titles & published) and key not in allow_missing:
                missing.append(key)
    return missing


def purge(cfg, *, deck_titles: list[str], podcast_ids: list[str], video_ids: list[str],
          record: Optional[Path], allow_missing: set[str], dry_run: bool = False) -> dict[str, Any]:
    from routers import decks as dr
    from routers import playlists as plr
    from routers import podcast as pr
    from routers import video as vr

    protected = {t for s in SEED_DECKS for t in s.titles}
    clash = sorted(set(deck_titles) & protected)
    if clash:
        raise MaintenanceError(f"refusing to delete new AI decks: {clash}")
    new_titles = {t for s in (*PODCASTS.values(), *VIDEOS.values()) for t in s.titles}

    inv = inventory(cfg)
    missing = missing_new_content(inv, allow_missing)
    if missing:
        raise MaintenanceError(f"new content not ready, nothing deleted: {missing}")

    decks = [d for d in inv["decks"] if d["title"] in set(deck_titles)]
    not_found = sorted(set(deck_titles) - {d["title"] for d in decks})
    deck_ids = {d["deck_id"] for d in decks}
    # Playlists are only touched by a deck purge; a media-only purge never deletes one.
    playlists = [p for p in inv["playlists"]
                 if p.get("status") == "broken" or any(s.get("deck_id") in deck_ids for s in p["steps"])
                 ] if deck_ids else []
    if any(p.get("title") == PLAYLIST_TITLE and p.get("status") != "broken" for p in playlists):
        raise MaintenanceError(f"{PLAYLIST_TITLE!r} references a deck slated for deletion")
    podcasts = [i for i in inv["podcasts"] if i.get("job_id") in set(podcast_ids)]
    videos = [i for i in inv["videos"] if i.get("job_id") in set(video_ids)]
    for items, doomed in ((inv["podcasts"], podcasts), (inv["videos"], videos)):
        doomed_ids = {i["job_id"] for i in doomed}
        for item in doomed:
            if item.get("title") not in new_titles:
                continue
            # A new-catalog item may only go once a published re-render with the same title remains.
            if not any(o.get("title") == item.get("title") and o.get("published")
                       and o.get("job_id") not in doomed_ids for o in items):
                raise MaintenanceError(f"refusing to delete new item {item.get('job_id')} ({item.get('title')})")
    not_found += sorted(set(podcast_ids) - {i["job_id"] for i in podcasts})
    not_found += sorted(set(video_ids) - {i["job_id"] for i in videos})

    result: dict[str, Any] = {
        "deleted_at": datetime.now(timezone.utc).isoformat(),
        "dry_run": dry_run,
        "playlists": [{"id": p["id"], "title": p.get("title"), "status": p.get("status"),
                       "steps": [s.get("deck_title") for s in p["steps"]]} for p in playlists],
        "decks": [{"deck_id": d["deck_id"], "title": d["title"], "language": d.get("language"),
                   "published": d.get("published"), "live_demo": d.get("live_demo")} for d in decks],
        "podcasts": [{"job_id": i["job_id"], "title": i.get("title"), "language": i.get("language"),
                      "speaker_names": i.get("speaker_names")} for i in podcasts],
        "videos": [{"job_id": i["job_id"], "title": i.get("title"), "language": i.get("language"),
                    "voice": i.get("voice")} for i in videos],
        "not_found": not_found,
        "errors": [],
    }
    if not dry_run:
        def _do(label: str, fn) -> None:
            try:
                fn()
                _say(f"deleted {label}")
            except Exception as exc:  # noqa: BLE001
                detail = getattr(exc, "detail", None) or str(exc)
                result["errors"].append({"item": label, "error": str(detail)[:300]})
                _say(f"FAILED {label}: {detail}")

        # Playlists first so the decks are no longer referenced (no forced delete).
        for p in playlists:
            _do(f"playlist {p['id']}", lambda p=p: plr.delete_playlist(p["id"], cfg=cfg))
        for d in decks:
            _do(f"deck {d['deck_id']} ({d['title']})", lambda d=d: dr.delete_deck(d["deck_id"], force=False, cfg=cfg))
        for i in podcasts:
            _do(f"podcast {i['job_id']}", lambda i=i: pr.delete_library_item(i["job_id"], cfg=cfg))
        for i in videos:
            _do(f"video {i['job_id']}", lambda i=i: vr.delete_library_item(i["job_id"], cfg=cfg))
    if record:
        record.parent.mkdir(parents=True, exist_ok=True)
        record.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


# ---------------------------------------------------------------------------
# Retitle (filename slug -> display title)
# ---------------------------------------------------------------------------
def _deck_search(cfg):
    from routers import decks as dr

    try:
        return dr.get_deck_search(cfg, dr.get_openai_client(cfg))
    except Exception:  # noqa: BLE001
        log.exception("deck search unavailable")
        return None


def retitle(cfg, *, dry_run: bool = False) -> dict[str, Any]:
    """Give slug-titled AI decks, podcasts and videos their display titles.

    Also rewrites the denormalized ``deck_title`` of playlist steps and of the
    deck's AI Search chunks, so search hits and playlists show the same title.
    Idempotent: items already carrying a display title are reported as
    ``already`` (their search chunks are re-synced, which is a no-op merge)."""
    from routers import decks as dr
    from routers import podcast as pr
    from routers import video as vr
    from services.storage import deck_title

    store = dr._store(cfg)
    if not store or not store.cosmos_available:
        raise MaintenanceError("Cosmos DB unavailable")
    display = set(DISPLAY_TITLES.values())
    report: dict[str, Any] = {"dry_run": dry_run, "decks": [], "playlists": [], "podcasts": [], "videos": [],
                              "errors": []}

    def _fail(item: str, msg: str) -> None:
        report["errors"].append({"item": item, "error": msg[:300]})
        _say(f"FAILED {item}: {msg}")

    final_titles: dict[str, str] = {}
    search = None if dry_run else _deck_search(cfg)
    for d in store.list_decks():
        if not isinstance(d, dict) or not d.get("id"):
            continue
        current = deck_title(d)
        target = DISPLAY_TITLES.get(current)
        if target is None and current not in display:
            continue
        entry: dict[str, Any] = {"deck_id": d["id"], "before": current, "after": target or current}
        if target is None:
            entry["action"] = "already"
        elif dry_run:
            entry["action"] = "would rename"
        elif store.set_deck_title(d["id"], target) is None:
            entry["action"] = "failed"
            _fail(f"deck {d['id']}", "title update failed")
        else:
            entry["action"] = "renamed"
        if entry["action"] != "failed":
            final_titles[d["id"]] = entry["after"]
            if search is not None:
                try:
                    entry["search_chunks"] = search.set_deck_title(d["id"], entry["after"])
                except Exception as exc:  # noqa: BLE001
                    _fail(f"search {d['id']}", str(exc))
        _say(f"deck {current!r} -> {entry['after']!r}: {entry['action']}")
        report["decks"].append(entry)

    for p in store.list_playlists():
        stale = [s for s in p.get("steps") or []
                 if s.get("deck_id") in final_titles and s.get("deck_title") != final_titles[s["deck_id"]]]
        if not stale:
            continue
        entry = {"id": p.get("id"), "title": p.get("title"),
                 "steps": [{"before": s.get("deck_title"), "after": final_titles[s["deck_id"]]} for s in stale]}
        if dry_run:
            entry["action"] = "would update"
        else:
            full = store.get_playlist(p["id"])
            if not full:
                entry["action"] = "failed"
                _fail(f"playlist {p.get('id')}", "not found")
            else:
                for s in full.get("steps") or []:
                    if s.get("deck_id") in final_titles:
                        s["deck_title"] = final_titles[s["deck_id"]]
                if store.save_playlist(full):
                    entry["action"] = "updated"
                else:
                    entry["action"] = "failed"
                    _fail(f"playlist {p.get('id')}", "save failed")
        _say(f"playlist {p.get('title')!r}: {entry['action']}")
        report["playlists"].append(entry)

    for kind, lib in (("podcasts", pr._get_library(cfg)), ("videos", vr._get_library(cfg))):
        if not lib.available:
            _fail(kind, "library storage unavailable")
            continue
        for item in lib.list():
            current = item.get("title") or ""
            target = DISPLAY_TITLES.get(current)
            doc_target = DISPLAY_TITLES.get(item.get("document_title") or "")
            if target is None and doc_target is None:
                if current in display:
                    report[kind].append({"job_id": item.get("job_id"), "before": current, "after": current,
                                         "action": "already"})
                continue
            entry = {"job_id": item.get("job_id"), "before": current, "after": target or current}
            if doc_target:
                entry["document_title"] = {"before": item.get("document_title"), "after": doc_target}
            if dry_run:
                entry["action"] = "would rename"
            elif lib.set_title(item["job_id"], target or current, document_title=doc_target) is None:
                entry["action"] = "failed"
                _fail(f"{kind[:-1]} {item.get('job_id')}", "title update failed")
            else:
                entry["action"] = "renamed"
            _say(f"{kind[:-1]} {item.get('job_id')} {current!r} -> {entry['after']!r}: {entry['action']}")
            report[kind].append(entry)
    return report


# ---------------------------------------------------------------------------
# Rebrand decks in place (same deck ids)
# ---------------------------------------------------------------------------
def _slide_text(content: bytes) -> list[dict[str, str]]:
    """Title/body/notes per slide, extracted exactly like the production parser."""
    from pptx import Presentation

    from services import pptx_parser as pp

    prs = Presentation(io.BytesIO(content))
    return [{"title": pp._extract_title(s), "body": pp._extract_body(s), "notes": pp._extract_notes(s)}
            for s in prs.slides]


def _versioned(url: str, version: int) -> str:
    return url if "?" in url else f"{url}?v={version}"


def rebrand_decks(cfg, *, dry_run: bool = False, deck_dir: Optional[Path] = None,
                  only: Optional[set[str]] = None) -> dict[str, Any]:
    """Overwrite the AI decks' slide images and slide text from the regenerated .pptx.

    Each existing deck keeps its id, title, speaker notes, translated notes,
    published/live-demo flags and playlist membership; a deck is skipped
    (error) when its slide count or any speaker note differs from the new file.
    Slide PNGs are overwritten under the same blob names, ``images_version`` is
    bumped (appended as ``?v=`` to unsigned image URLs) and the AI Search chunks
    are re-indexed with the new slide text. Idempotent: a deck whose stored
    ``source_sha256`` and slide text already match is reported ``already``.
    Never deletes decks, playlists or media."""
    from routers import decks as dr
    from services.deck_search import DeckSlide
    from services.pptx_parser import _fallback_placeholders, parse_pptx
    from services.storage import deck_title

    store = dr._store(cfg)
    if not store or not store.cosmos_available:
        raise MaintenanceError("Cosmos DB unavailable")
    if not dry_run and not store.available:
        raise MaintenanceError("Blob Storage unavailable")
    by_title: dict[str, dict] = {}
    for d in store.list_decks():
        if isinstance(d, dict) and d.get("id"):
            by_title.setdefault(dr._doc_to_summary(d).title, d)

    report: dict[str, Any] = {"dry_run": dry_run, "decks": [], "errors": []}
    search = None
    for spec in SEED_DECKS:
        if only and spec.slug not in only:
            continue
        entry: dict[str, Any] = {"title": spec.title, "file": spec.file}
        report["decks"].append(entry)

        def _fail(msg: str, entry=entry, spec=spec) -> None:
            entry["action"] = "aborted"
            entry["error"] = msg
            report["errors"].append({"item": f"deck {spec.slug}", "error": msg[:300]})
            _say(f"FAILED deck {spec.title}: {msg}")

        found = by_title.get(spec.title) or by_title.get(spec.slug)
        if not found:
            _fail("deck not found (run seed first)")
            continue
        deck_id = found["id"]
        entry["deck_id"] = deck_id
        doc = store.get_presentation(deck_id)
        if not doc or doc.get("source") != "deck":
            _fail("deck document missing")
            continue
        content = _deck_path(spec.file, deck_dir).read_bytes()
        if content[:4] != b"PK\x03\x04":
            raise MaintenanceError(f"{spec.file} is not a .pptx")
        sha = hashlib.sha256(content).hexdigest()
        new = _slide_text(content)
        old = sorted((s for s in doc.get("slides") or [] if isinstance(s, dict)),
                     key=lambda s: int(s.get("index", 0)))
        if len(new) != len(old) or int(doc.get("slide_count") or len(old)) != len(new):
            _fail(f"slide count differs (stored {len(old)}, file {len(new)})")
            continue
        bad = [i for i, (o, n) in enumerate(zip(old, new)) if (o.get("notes") or "") != n["notes"]]
        if bad:
            _fail(f"speaker notes differ on slides {[i + 1 for i in bad]}")
            continue
        changed = [i for i, (o, n) in enumerate(zip(old, new))
                   if (o.get("body") or "") != n["body"] or (o.get("title") or "") != n["title"]]
        entry["text_changes"] = len(changed)
        if doc.get("source_sha256") == sha and not changed:
            entry["action"] = "already"
            _say(f"deck {spec.title}: already rebranded")
            continue
        if dry_run:
            entry["action"] = "would rebrand"
            _say(f"deck {spec.title}: would rebrand ({len(new)} slides)")
            continue

        _say(f"deck {spec.title}: render {spec.file}")
        pres = parse_pptx(content, spec.file, cfg.libreoffice_path or "soffice")
        images = dict(pres.slide_images)
        if sorted(images) != list(range(len(new))):
            _fail(f"rendered {len(images)} images for {len(new)} slides")
            continue
        placeholders = dict(_fallback_placeholders(len(new)))
        if any(images[i] == placeholders[i] for i in images):
            _fail("slide rendering fell back to placeholders (LibreOffice unavailable?)")
            continue
        if [s.notes for s in pres.slides] != [n["notes"] for n in new]:
            _fail("parser notes differ from the text extraction")
            continue

        local_dir = dr.SLIDES_DIR / deck_id
        uploaded = 0
        for idx in sorted(images):
            if not store.upload_slide_image(deck_id, idx, images[idx]):
                break
            uploaded += 1
            local_dir.mkdir(parents=True, exist_ok=True)
            (local_dir / f"{idx}.png").write_bytes(images[idx])
        entry["images"] = uploaded
        if uploaded != len(images):
            _fail(f"blob upload failed after {uploaded}/{len(images)} images")
            continue
        if doc.get("pptx_url"):
            entry["pptx"] = "replaced" if store.upload_pptx(deck_id, doc.get("filename") or spec.file,
                                                            content) else "failed"

        latest = store.get_presentation(deck_id) or doc
        version = int(latest.get("images_version") or 0) + 1
        for s in latest.get("slides") or []:
            if not isinstance(s, dict):
                continue
            i = int(s.get("index", 0))
            s["title"], s["body"] = pres.slides[i].title, pres.slides[i].body
            url = store.get_slide_image_url(deck_id, i) or f"/api/slides/{deck_id}/{i}.png"
            s["image_url"] = _versioned(url, version)
        latest["images_version"] = version
        latest["source_sha256"] = sha
        latest["rebranded_at"] = datetime.now(timezone.utc).isoformat()
        if not store.save_presentation(latest):
            _fail("Cosmos save failed (images already replaced)")
            continue
        entry.update(action="rebranded", images_version=version)

        if search is None:
            search = _deck_search(cfg) or False
        if search:
            try:
                entry["search_chunks"] = search.index_deck([
                    DeckSlide(deck_id=deck_id, deck_title=deck_title(latest), slide_index=i,
                              slide_title=s.title or f"Slide {i + 1}", body=s.body or "", notes=s.notes or "",
                              language=latest.get("language") or spec.language)
                    for i, s in enumerate(pres.slides)
                ])
            except Exception as exc:  # noqa: BLE001
                report["errors"].append({"item": f"search {deck_id}", "error": str(exc)[:300]})
                _say(f"FAILED search reindex {deck_id}: {exc}")
        _say(f"deck {spec.title}: rebranded ({uploaded} images, v{version})")
    return report


# ---------------------------------------------------------------------------
# Everything (one operator call)
# ---------------------------------------------------------------------------
async def publish_all(cfg, *, dry_run: bool = False, deck_dir: Optional[Path] = None,
                      skip: frozenset[str] | set[str] = frozenset()) -> dict[str, Any]:
    """Seed the decks, then render each podcast and video in turn.

    Renders run one at a time (batch avatar synthesis is throttled per
    resource); a failed item is reported and the next one still runs."""
    summary: dict[str, Any] = {"seed": None, "podcasts": [], "videos": [], "errors": []}
    try:
        summary["seed"] = await seed(cfg, dry_run=dry_run, deck_dir=deck_dir)
    except Exception as exc:  # noqa: BLE001
        log.exception("seed failed")
        summary["errors"].append({"item": "seed", "error": str(getattr(exc, "detail", exc))[:300]})
    for kind, names, fn in (("podcast", PODCASTS, render_podcast), ("video", VIDEOS, render_video)):
        for name in names:
            key = f"{kind}:{name}"
            if key in skip:
                summary[f"{kind}s"].append({"name": name, "action": "skipped"})
                continue
            try:
                summary[f"{kind}s"].append(await fn(cfg, name, dry_run=dry_run, deck_dir=deck_dir))
            except Exception as exc:  # noqa: BLE001
                log.exception("%s failed", key)
                summary["errors"].append({"item": key, "error": str(getattr(exc, "detail", exc))[:300]})
                _say(f"FAILED {key}: {exc}")
    return summary


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def _print(obj: Any) -> None:
    print(json.dumps(obj, indent=2, ensure_ascii=False, default=str), flush=True)


def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(prog="python -m maintenance.content", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--deck-dir", type=Path, default=None, help="Directory holding the .pptx decks")
    ap.add_argument("--background", type=Path, default=None, metavar="LOG",
                    help="Detach (new session) and append all output to LOG; poll it with tail")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("inventory", help="List decks, playlists, podcasts and videos (no URLs)")
    p.add_argument("--json", type=Path, default=None)

    p = sub.add_parser("seed", help="Upload/publish/translate the AI decks and create the playlist")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("render-podcast", help="Render + publish an AI podcast")
    p.add_argument("name", choices=[*PODCASTS, "all"])
    p.add_argument("--force", action="store_true",
                   help="Render a new item even if one with this title exists (purge the old id afterwards)")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("render-video", help="Render + publish an AI video")
    p.add_argument("name", choices=[*VIDEOS, "all"])
    p.add_argument("--force", action="store_true",
                   help="Render a new item even if one with this title exists (purge the old id afterwards)")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("purge", help="Delete retired decks/podcasts/videos (+ their playlists)")
    p.add_argument("--deck-title", action="append", default=[])
    p.add_argument("--podcast-id", action="append", default=[])
    p.add_argument("--video-id", action="append", default=[])
    p.add_argument("--record", type=Path, default=None, help="Write the deletion record (JSON) here")
    p.add_argument("--allow-missing", action="append", default=[],
                   help="New item allowed to be absent, e.g. video:practitioner")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("retitle", help="Rename slug-titled AI decks/podcasts/videos to display titles")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("rebrand-decks",
                       help="Re-render the AI decks' slide images/text in place (same deck ids, notes unchanged)")
    p.add_argument("--deck", action="append", default=[], choices=[s.slug for s in SEED_DECKS],
                   help="Limit to these decks (default: all AI decks)")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("all", help="seed, then every podcast, then every video (continues past failures)")
    p.add_argument("--skip", action="append", default=[], help="e.g. video:practitioner")
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("logtail", help="Print the end of a log with URL query strings redacted")
    p.add_argument("log", type=Path)
    p.add_argument("-n", type=int, default=40)
    return ap


_QUERY_RE = None


def redact(line: str) -> str:
    """Drop URL query strings (SAS tokens) and key-like values from a log line."""
    import re

    global _QUERY_RE
    if _QUERY_RE is None:
        _QUERY_RE = re.compile(r"(https?://[^\s?'\"]+)\?[^\s'\"]*")
    line = _QUERY_RE.sub(r"\1?<redacted>", line)
    return re.sub(r"(?i)((?:sig|key|token|secret|password)=)[^\s&'\"]+", r"\1<redacted>", line)


def logtail(path: Path, n: int) -> list[str]:
    if not path.is_file():
        return [f"(no log at {path})"]
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()[-max(1, n):]
    return [redact(line) for line in lines]


def _detach(log_path: Path) -> None:
    """Double-fork into a new session writing to ``log_path`` (POSIX only).

    Lets an operator start a long render through a short-lived
    ``az containerapp exec`` call and poll the log afterwards."""
    import os

    log_path.parent.mkdir(parents=True, exist_ok=True)
    pid = os.fork()
    if pid > 0:
        print(f"started in background (pid {pid}); log: {log_path}", flush=True)
        os._exit(0)
    os.setsid()
    if os.fork() > 0:
        os._exit(0)
    fd = os.open(str(log_path), os.O_WRONLY | os.O_CREAT | os.O_APPEND, 0o600)
    devnull = os.open(os.devnull, os.O_RDONLY)
    os.dup2(devnull, 0)
    os.dup2(fd, 1)
    os.dup2(fd, 2)


def main(argv: Optional[list[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    if args.cmd == "logtail":
        print("\n".join(logtail(args.log, args.n)), flush=True)
        return 0
    if args.background:
        _detach(args.background)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    for noisy in ("azure", "httpx", "openai", "urllib3"):
        logging.getLogger(noisy).setLevel(logging.WARNING)
    cfg = _cfg()
    try:
        if args.cmd == "inventory":
            inv = inventory(cfg)
            if args.json:
                args.json.write_text(json.dumps(inv, indent=2, ensure_ascii=False), encoding="utf-8")
            _print(inv)
        elif args.cmd == "seed":
            _print(asyncio.run(seed(cfg, dry_run=args.dry_run, deck_dir=args.deck_dir)))
        elif args.cmd == "render-podcast":
            names = list(PODCASTS) if args.name == "all" else [args.name]
            _print([asyncio.run(render_podcast(cfg, n, dry_run=args.dry_run, deck_dir=args.deck_dir,
                                               force=args.force))
                    for n in names])
        elif args.cmd == "render-video":
            names = list(VIDEOS) if args.name == "all" else [args.name]
            _print([asyncio.run(render_video(cfg, n, dry_run=args.dry_run, deck_dir=args.deck_dir,
                                             force=args.force))
                    for n in names])
        elif args.cmd == "purge":
            result = purge(cfg, deck_titles=args.deck_title, podcast_ids=args.podcast_id,
                           video_ids=args.video_id, record=args.record,
                           allow_missing=set(args.allow_missing), dry_run=args.dry_run)
            _print(result)
            if result["errors"]:
                return 1
        elif args.cmd == "retitle":
            result = retitle(cfg, dry_run=args.dry_run)
            _print(result)
            if result["errors"]:
                return 1
        elif args.cmd == "rebrand-decks":
            result = rebrand_decks(cfg, dry_run=args.dry_run, deck_dir=args.deck_dir,
                                   only=set(args.deck) or None)
            _print(result)
            if result["errors"]:
                return 1
        elif args.cmd == "all":
            summary = asyncio.run(publish_all(cfg, dry_run=args.dry_run, deck_dir=args.deck_dir,
                                              skip=set(args.skip)))
            _print(summary)
            if summary["errors"]:
                return 1
    except MaintenanceError as exc:
        _say(f"ERROR: {exc}")
        return 2
    _say("done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
